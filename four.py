"""2026 CUMCM C题问题四：傅里叶—小波实时电价预测下的 LP 与 MPC 购电策略。

电价预测链：
1. 严格前向贝叶斯优化历史窗口、权重及分解阶数，并采用早停；
2. 对加权电价轮廓作傅里叶拟合，提取日周期；
3. 对拟合残差作Haar小波软阈值重构，恢复局部尖峰与突变；
4. 使用前一日同一时刻的预测偏差滚动修正，并在线估计修正系数；
5. 对最终预测作非负截断。

问题4-2沿用问题二的48小时滚动线性规划；问题4-3沿用问题三的MPC-LP，
分别生成官方模板要求的result4-2.xlsx和result4-3.xlsx。
问题4-3策略只按验证集费用选择，测试集不参与调参或策略选择。

依赖：two.py、three.py、numpy、scipy、openpyxl。
正式求解：python four.py
贝叶斯寻优：python four.py --bayes-optimize
无Hampel消融：python four.py --bayes-optimize --bayes-disable-hampel
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import os
import shutil
from copy import copy
from dataclasses import asdict, dataclass, replace
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

import three
import two


ROOT = Path(__file__).resolve().parent
DEFAULT_PRICE = ROOT / "C题" / "附件" / "附件4.xlsx"
DEFAULT_FORECAST = ROOT / "C题" / "附件" / "csv" / "附件3.csv"
DEFAULT_TEMPLATE_42 = ROOT / "C题" / "附件" / "附件5" / "result4-2.xlsx"
DEFAULT_TEMPLATE_43 = ROOT / "C题" / "附件" / "附件5" / "result4-3.xlsx"
DEFAULT_OUTPUT_42 = ROOT / "C题" / "result4-2.xlsx"
DEFAULT_OUTPUT_43 = ROOT / "C题" / "result4-3.xlsx"
DEFAULT_SUMMARY_42 = ROOT / "C题" / "问题四_对应问题二逐日汇总.csv"
DEFAULT_SUMMARY_43 = ROOT / "C题" / "问题四_对应问题三逐日汇总.csv"
DEFAULT_COMPARISON = ROOT / "C题" / "问题四_MPC策略对比.csv"
DEFAULT_PRICE_EVALUATION = ROOT / "C题" / "问题四_电价预测评估.csv"
DEFAULT_SPLIT_REPORT = ROOT / "C题" / "问题四_训练验证测试检验.csv"
DEFAULT_MODEL = ROOT / "C题" / "问题四_动态电价数学模型.md"
DEFAULT_PRICE_BAYES_REPORT = ROOT / "C题" / "贝叶斯优化_Hampel电价参数与评价_严格前向.json"
DEFAULT_PRICE_BAYES_CSV = ROOT / "C题" / "贝叶斯优化_Hampel电价逐日评价_严格前向.csv"
DEFAULT_PRICE_BAYES_NO_HAMPEL_REPORT = ROOT / "C题" / "贝叶斯优化_无Hampel电价参数与评价_严格前向.json"
DEFAULT_PRICE_BAYES_NO_HAMPEL_CSV = ROOT / "C题" / "贝叶斯优化_无Hampel电价逐日评价_严格前向.csv"


@dataclass(frozen=True)
class Config(three.Config):
    price_history_days: int = 28
    price_half_life_days: float = 7.0
    price_same_weekday_multiplier: float = 2.0
    price_fourier_harmonics: int = 6
    price_wavelet_level: int = 3
    price_bias_days: int = 28
    minimum_bias_pairs: int = 7
    default_bias_weight: float = 0.60
    use_bayesian_price_schedule: bool = True


# 严格前向贝叶斯优化得到的冻结电价参数时间表。每次切换均从验证日的
# 下一日开始生效，保证当日真实价格不会参与当日预测。
_PRICE_PARAMETER_SCHEDULE = (
    (
        date(2025, 2, 28),
        dict(
            price_history_days=40,
            price_half_life_days=28.260560329514153,
            price_same_weekday_multiplier=3.335679870659064,
            price_fourier_harmonics=5,
            price_wavelet_level=2,
            price_bias_days=32,
            default_bias_weight=0.6296739771158961,
        ),
    ),
    (
        date(2025, 3, 6),
        dict(
            price_history_days=31,
            price_half_life_days=24.43488696871456,
            price_same_weekday_multiplier=3.46299237674132,
            price_fourier_harmonics=5,
            price_wavelet_level=2,
            price_bias_days=47,
            default_bias_weight=0.6456718349291368,
        ),
    ),
    (
        date(2025, 3, 12),
        dict(
            price_history_days=33,
            price_half_life_days=25.145760809373442,
            price_same_weekday_multiplier=4.009445223594914,
            price_fourier_harmonics=3,
            price_wavelet_level=2,
            price_bias_days=38,
            default_bias_weight=0.44690879842079867,
        ),
    ),
    (
        date(2025, 4, 5),
        dict(
            price_history_days=22,
            price_half_life_days=27.116107451115006,
            price_same_weekday_multiplier=4.444737582066487,
            price_fourier_harmonics=2,
            price_wavelet_level=2,
            price_bias_days=51,
            default_bias_weight=0.6237624262745828,
        ),
    ),
)


def bayesian_price_config(cfg: Config, issue_date: date) -> Config:
    """返回预测发布日可用的已冻结电价参数，不读取未来数据。"""
    if not cfg.use_bayesian_price_schedule:
        return cfg
    selected = cfg
    for effective_date, parameters in _PRICE_PARAMETER_SCHEDULE:
        if issue_date < effective_date:
            break
        selected = replace(cfg, **parameters)
    return selected


@dataclass
class PriceData:
    dates: list[date]
    price: np.ndarray


@dataclass
class PriceForecast:
    value: np.ndarray
    periodic: np.ndarray
    wavelet_residual: np.ndarray
    previous_day_bias: np.ndarray
    bias_weight: float


@dataclass
class Daily42:
    day: date
    forecast_price: np.ndarray
    grid_kwh: np.ndarray
    charge_kwh: np.ndarray
    discharge_kwh: np.ndarray
    emergency_kwh: np.ndarray
    soc_start_kwh: float
    soc_end_kwh: float
    plan_cost_yuan: float
    emergency_cost_yuan: float

    @property
    def total_cost_yuan(self) -> float:
        return self.plan_cost_yuan + self.emergency_cost_yuan


def read_dynamic_prices(path: Path, chronological_order: list[int]) -> PriceData:
    if not path.exists():
        raise FileNotFoundError(f"找不到附件4：{path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    if worksheet.max_row != 366 or worksheet.max_column != 145:
        raise ValueError("附件4应包含1行表头、365天和每天144个实时电价。")

    dates: list[date] = []
    rows: list[list[float]] = []
    order = np.asarray(chronological_order, dtype=int)
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=2, max_row=366, values_only=True), start=2
    ):
        raw_date = row[0]
        if isinstance(raw_date, datetime):
            current_date = raw_date.date()
        elif isinstance(raw_date, date):
            current_date = raw_date
        else:
            current_date = datetime.fromisoformat(str(raw_date)).date()
        values = np.asarray(
            [two._float(value, f"附件4第{row_number}行电价") for value in row[1:]],
            dtype=float,
        )[order]
        dates.append(current_date)
        rows.append(values.tolist())
    workbook.close()

    expected = [date(2025, 1, 1) + timedelta(days=i) for i in range(365)]
    if dates != expected:
        raise ValueError("附件4日期应连续覆盖2025-01-01至2025-12-31。")
    matrix = np.asarray(rows, dtype=float)
    if not np.all(np.isfinite(matrix)) or np.any(matrix < 0):
        raise ValueError("附件4包含负值或非有限电价。")
    return PriceData(dates, matrix)


def _price_fourier_design(hours: np.ndarray, harmonics: int) -> np.ndarray:
    omega = 2.0 * np.pi / 24.0
    columns = [np.ones_like(hours)]
    for order in range(1, harmonics + 1):
        columns.extend(
            [np.cos(order * omega * hours), np.sin(order * omega * hours)]
        )
    return np.column_stack(columns)


def fit_price_fourier(
    hours: np.ndarray, profile: np.ndarray, harmonics: int
) -> np.ndarray:
    design = _price_fourier_design(hours, harmonics)
    coefficients, *_ = np.linalg.lstsq(design, profile, rcond=None)
    return design @ coefficients


def _price_history_profile(
    prices: PriceData,
    cutoff: int,
    target_date: date,
    baseline_price: np.ndarray,
    cfg: Config,
    allowed_indices: set[int] | None = None,
) -> np.ndarray:
    """只使用索引小于cutoff的已观测价格构造目标日轮廓。"""
    if cutoff <= 0:
        return baseline_price.copy()
    first = max(0, cutoff - cfg.price_history_days)
    indices = np.arange(first, cutoff, dtype=int)
    if allowed_indices is not None:
        indices = np.asarray([i for i in indices if i in allowed_indices], dtype=int)
    if indices.size == 0:
        return baseline_price.copy()
    ages = cutoff - 1 - indices
    weights = np.exp(-math.log(2.0) * ages / cfg.price_half_life_days)
    same_weekday = np.asarray(
        [prices.dates[index].weekday() == target_date.weekday() for index in indices],
        dtype=float,
    )
    weights *= 1.0 + cfg.price_same_weekday_multiplier * same_weekday
    return np.average(prices.price[indices], axis=0, weights=weights)


def _causal_price_bases(
    prices: PriceData,
    baseline_price: np.ndarray,
    cfg: Config,
    training_cutoff: int | None = None,
    training_indices: set[int] | None = None,
) -> tuple[list[np.ndarray], list[np.ndarray], list[np.ndarray]]:
    hours = np.arange(144, dtype=float) * cfg.dt_hours
    bases: list[np.ndarray] = []
    periodic_parts: list[np.ndarray] = []
    wavelet_parts: list[np.ndarray] = []
    for cutoff, current_date in enumerate(prices.dates):
        model_cutoff = (
            cutoff
            if training_cutoff is None
            else min(cutoff, training_cutoff)
        )
        allowed_indices = training_indices
        profile = _price_history_profile(
            prices, model_cutoff, current_date, baseline_price, cfg, allowed_indices
        )
        periodic = fit_price_fourier(
            hours, profile, cfg.price_fourier_harmonics
        )
        wavelet = two.wavelet_correct_residual(
            profile, periodic, cfg.price_wavelet_level
        )
        bases.append(periodic + wavelet)
        periodic_parts.append(periodic)
        wavelet_parts.append(wavelet)
    return bases, periodic_parts, wavelet_parts


def _estimate_bias_weight(
    prices: PriceData,
    causal_bases: list[np.ndarray],
    cutoff: int,
    cfg: Config,
    allowed_indices: set[int] | None = None,
) -> float:
    """用此前日期残差的一阶关系估计前一日偏差修正系数。"""
    pair_end = cutoff - 1
    pair_start = max(1, pair_end - cfg.price_bias_days + 1)
    if pair_end - pair_start + 1 < cfg.minimum_bias_pairs:
        return cfg.default_bias_weight
    previous: list[np.ndarray] = []
    current: list[np.ndarray] = []
    for index in range(pair_start, pair_end + 1):
        if allowed_indices is not None and (index not in allowed_indices or index - 1 not in allowed_indices):
            continue
        previous.append(prices.price[index - 1] - causal_bases[index - 1])
        current.append(prices.price[index] - causal_bases[index])
    if len(previous) < cfg.minimum_bias_pairs:
        return cfg.default_bias_weight
    x = np.concatenate(previous)
    y = np.concatenate(current)
    denominator = float(np.dot(x, x))
    if denominator <= 1.0e-12:
        return cfg.default_bias_weight
    return float(np.clip(np.dot(x, y) / denominator, 0.0, 1.0))


def make_price_forecast(
    prices: PriceData,
    baseline_price: np.ndarray,
    causal_bases: list[np.ndarray],
    cutoff: int,
    target_date: date,
    cfg: Config,
    training_cutoff: int | None = None,
    training_indices: set[int] | None = None,
) -> PriceForecast:
    hours = np.arange(144, dtype=float) * cfg.dt_hours
    model_cutoff = (
        cutoff
        if training_cutoff is None
        else min(cutoff, training_cutoff)
    )
    profile = _price_history_profile(
        prices, model_cutoff, target_date, baseline_price, cfg, training_indices
    )
    periodic = fit_price_fourier(hours, profile, cfg.price_fourier_harmonics)
    wavelet = two.wavelet_correct_residual(
        profile, periodic, cfg.price_wavelet_level
    )
    latest_observed = cutoff - 1
    if latest_observed >= 0 and (training_indices is None or latest_observed in training_indices):
        bias = prices.price[latest_observed] - causal_bases[latest_observed]
        # 偏差修正系数只在训练集估计；前一日偏差仍是预测时可得信息。
        weight = _estimate_bias_weight(
            prices, causal_bases, model_cutoff, cfg, training_indices
        )
    else:
        bias = np.zeros(144, dtype=float)
        weight = 0.0
    value = np.maximum(0.0, periodic + wavelet + weight * bias)
    return PriceForecast(value, periodic, wavelet, bias, weight)


def prepare_daily_price_forecasts(
    prices: PriceData,
    baseline_price: np.ndarray,
    cfg: Config,
    training_cutoff: int | None = None,
    training_indices: set[int] | None = None,
) -> tuple[list[PriceForecast], list[np.ndarray]]:
    daily_configs = [bayesian_price_config(cfg, day) for day in prices.dates]
    bases_cache: dict[Config, list[np.ndarray]] = {}
    for forecast_cfg in dict.fromkeys(daily_configs):
        bases_cache[forecast_cfg], _, _ = _causal_price_bases(
            prices,
            baseline_price,
            forecast_cfg,
            training_cutoff=training_cutoff,
            training_indices=training_indices,
        )
    current: list[PriceForecast] = []
    following: list[np.ndarray] = []
    for cutoff, (current_date, forecast_cfg) in enumerate(
        zip(prices.dates, daily_configs)
    ):
        causal_bases = bases_cache[forecast_cfg]
        current.append(
            make_price_forecast(
                prices,
                baseline_price,
                causal_bases,
                cutoff,
                current_date,
                forecast_cfg,
                training_cutoff=training_cutoff,
                training_indices=training_indices,
            )
        )
        following.append(
            make_price_forecast(
                prices,
                baseline_price,
                causal_bases,
                cutoff,
                current_date + timedelta(days=1),
                forecast_cfg,
                training_cutoff=training_cutoff,
                training_indices=training_indices,
            ).value
        )
    return current, following


def simulate_42(
    baseline: two.BaselineData,
    historical: two.HistoricalData,
    prices: PriceData,
    current_prices: list[PriceForecast],
    following_prices: list[np.ndarray],
    cfg: Config,
    training_indices: set[int] | None = None,
) -> list[Daily42]:
    hours = np.arange(144, dtype=float) * cfg.dt_hours
    error_history: list[np.ndarray] = []
    results: list[Daily42] = []
    soc = cfg.initial_soc_kwh

    for day_index, current_date in enumerate(historical.dates):
        allowed_indices = training_indices
        forecast_cfg = two.bayesian_load_config(cfg, current_date)
        current_forecast = two.make_forecast(
            historical, baseline, day_index, current_date, hours, forecast_cfg,
            allowed_indices=allowed_indices,
            pv_cfg=cfg,
        )
        next_forecast = two.make_forecast(
            historical,
            baseline,
            day_index,
            current_date + timedelta(days=1),
            hours,
            forecast_cfg,
            allowed_indices=allowed_indices,
            pv_cfg=cfg,
        )
        reserve = two.risk_reserve(error_history, cfg)
        price_horizon = np.concatenate(
            [current_prices[day_index].value, following_prices[day_index]]
        )
        solution = two.solve_rolling_lp(
            price_horizon,
            np.concatenate(
                [
                    current_forecast.load_kw + reserve / cfg.dt_hours,
                    next_forecast.load_kw + reserve / cfg.dt_hours,
                ]
            ),
            np.concatenate([current_forecast.pv_kw, next_forecast.pv_kw]),
            soc,
            cfg,
        )
        grid = solution.grid_kwh[:144]
        charge = solution.charge_kwh[:144]
        discharge = solution.discharge_kwh[:144]
        emergency, _ = two.replay_actual_day(
            historical.load_kw[day_index],
            historical.pv_kw[day_index],
            grid,
            charge,
            discharge,
            cfg,
        )
        soc_start = soc
        soc = float(solution.soc_kwh[144])
        actual_price = prices.price[day_index]
        plan_cost = float(np.dot(actual_price, grid))
        emergency_cost = float(
            np.dot(cfg.emergency_price_multiplier * actual_price, emergency)
        )
        results.append(
            Daily42(
                current_date,
                current_prices[day_index].value,
                grid,
                charge,
                discharge,
                emergency,
                soc_start,
                soc,
                plan_cost,
                emergency_cost,
            )
        )
        actual_net = (
            historical.load_kw[day_index] - historical.pv_kw[day_index]
        ) * cfg.dt_hours
        forecast_net = (
            current_forecast.load_kw - current_forecast.pv_kw
        ) * cfg.dt_hours
        if training_indices is None or day_index in training_indices:
            error_history.append(actual_net - forecast_net)
        if (day_index + 1) % 60 == 0 or day_index == 364:
            print(f"问题4-2已完成{current_date}，SOC={soc:.2f} kWh")
    return results


def price_horizon(
    day_index: int,
    release_hour: int,
    current_prices: list[PriceForecast],
    following_prices: list[np.ndarray],
) -> np.ndarray:
    start = release_hour * 6
    if start == 0:
        return current_prices[day_index].value.copy()
    return np.concatenate(
        [current_prices[day_index].value[start:], following_prices[day_index][:start]]
    )


def simulate_43_strategy(
    name: str,
    release_hours: tuple[int, ...],
    historical: two.HistoricalData,
    prices: PriceData,
    decision_inputs: dict[tuple[int, int], three.DecisionInput],
    current_prices: list[PriceForecast],
    following_prices: list[np.ndarray],
    cfg: Config,
) -> three.StrategySummary:
    results: list[three.DailyResult] = []
    soc = cfg.initial_soc_kwh
    for day_index, current_date in enumerate(historical.dates):
        soc_start = soc
        decision_zero = decision_inputs[(day_index, 0)]
        plan_solution = two.solve_rolling_lp(
            price_horizon(
                day_index, 0, current_prices, following_prices
            ),
            decision_zero.load_kw + decision_zero.reserve_kwh / cfg.dt_hours,
            decision_zero.pv_kw,
            soc,
            cfg,
        )
        plan_grid = plan_solution.grid_kwh.copy()
        final_grid = plan_solution.grid_kwh.copy()
        final_charge = plan_solution.charge_kwh.copy()
        final_discharge = plan_solution.discharge_kwh.copy()

        for release_hour in (0, 6, 12, 18):
            start = release_hour * 6
            if release_hour > 0 and release_hour in release_hours:
                decision = decision_inputs[(day_index, release_hour)]
                adjusted = three.solve_adjustment_lp(
                    price_horizon(
                        day_index,
                        release_hour,
                        current_prices,
                        following_prices,
                    ),
                    decision.load_kw + decision.reserve_kwh / cfg.dt_hours,
                    decision.pv_kw,
                    soc,
                    plan_grid[start:],
                    144 - start,
                    cfg,
                )
                final_grid[start:] = adjusted.grid_kwh[: 144 - start]
                final_charge[start:] = adjusted.charge_kwh[: 144 - start]
                final_discharge[start:] = adjusted.discharge_kwh[: 144 - start]
            stop = min(start + 36, 144)
            soc = three.execute_soc(
                soc,
                final_charge[start:stop],
                final_discharge[start:stop],
                cfg,
            )

        emergency, curtailment = two.replay_actual_day(
            historical.load_kw[day_index],
            historical.pv_kw[day_index],
            final_grid,
            final_charge,
            final_discharge,
            cfg,
        )
        actual_price = prices.price[day_index]
        plan_cost = float(np.dot(actual_price, plan_grid))
        settled_cost = three.settled_purchase_cost(
            actual_price, plan_grid, final_grid, cfg
        )
        emergency_cost = float(
            np.dot(cfg.emergency_price_multiplier * actual_price, emergency)
        )
        results.append(
            three.DailyResult(
                current_date,
                plan_grid,
                final_grid,
                final_charge,
                final_discharge,
                emergency,
                curtailment,
                soc_start,
                soc,
                plan_cost,
                settled_cost,
                emergency_cost,
            )
        )
        if (day_index + 1) % 120 == 0 or day_index == 364:
            print(f"{name}已完成{current_date}，SOC={soc:.2f} kWh")

    selected = [item for item in results if item.day >= date(2025, 2, 1)]
    plan_energy = float(sum(np.sum(item.plan_grid_kwh) for item in selected))
    final_energy = float(sum(np.sum(item.final_grid_kwh) for item in selected))
    increase_energy = float(
        sum(
            np.sum(np.maximum(0.0, item.final_grid_kwh - item.plan_grid_kwh))
            for item in selected
        )
    )
    decrease_energy = float(
        sum(
            np.sum(np.maximum(0.0, item.plan_grid_kwh - item.final_grid_kwh))
            for item in selected
        )
    )
    emergency_energy = float(sum(np.sum(item.emergency_kwh) for item in selected))
    plan_cost = float(sum(item.plan_cost_yuan for item in selected))
    settled_cost = float(sum(item.settled_purchase_cost_yuan for item in selected))
    emergency_cost = float(sum(item.emergency_cost_yuan for item in selected))
    return three.StrategySummary(
        name,
        release_hours,
        results,
        plan_energy,
        final_energy,
        increase_energy,
        decrease_energy,
        emergency_energy,
        plan_cost,
        settled_cost,
        emergency_cost,
        settled_cost + emergency_cost,
    )


def validate_42(results: list[Daily42], cfg: Config) -> None:
    if len(results) != 365:
        raise RuntimeError("问题4-2没有生成365天结果。")
    for index, item in enumerate(results):
        if index and abs(item.soc_start_kwh - results[index - 1].soc_end_kwh) > 1e-4:
            raise RuntimeError(f"问题4-2在{item.day}的跨日SOC不连续。")
        if np.any(item.grid_kwh < -1e-6) or np.any(item.emergency_kwh < -1e-6):
            raise RuntimeError(f"问题4-2在{item.day}出现负购电量。")
        if not cfg.soc_min_kwh - 1e-4 <= item.soc_end_kwh <= cfg.soc_max_kwh + 1e-4:
            raise RuntimeError(f"问题4-2在{item.day}的SOC越界。")


def _storage_rows(results: list, q43: bool = False) -> list[list[object]]:
    rows: list[list[object]] = []
    periods = [
        "0:00-4:00",
        "4:00-8:00",
        "8:00-12:00",
        "12:00-16:00",
        "16:00-20:00",
        "20:00-24:00",
    ]
    for item in results:
        charge = item.charge_kwh
        discharge = item.discharge_kwh
        for block, period in enumerate(periods):
            start, stop = block * 24, (block + 1) * 24
            rows.append(
                [
                    item.day.isoformat() if block == 0 else None,
                    period,
                    round(float(np.sum(charge[start:stop])), 4),
                    round(float(np.sum(discharge[start:stop])), 4),
                    "0:00" if block == 0 else ("24:00" if block == 1 else None),
                    round(
                        float(item.soc_start_kwh if block == 0 else item.soc_end_kwh),
                        4,
                    )
                    if block in (0, 1)
                    else None,
                ]
            )
    return rows


def _emergency_rows(results: list, tolerance: float) -> list[list[object]]:
    rows: list[list[object]] = []
    for item in results:
        groups = two.group_emergency(item.emergency_kwh, tolerance)
        for group_index, (period, amount) in enumerate(groups):
            rows.append(
                [
                    item.day.isoformat() if group_index == 0 else None,
                    period,
                    round(amount, 4),
                ]
            )
    return rows or [["无紧急购电", None, 0.0]]


def make_workbook_payload(
    results42: list[Daily42],
    selected43: three.StrategySummary,
    chronological_order: list[int],
    cfg: Config,
) -> dict[str, object]:
    selected42 = [item for item in results42 if item.day >= date(2025, 2, 1)]
    selected43_results = [
        item for item in selected43.results if item.day >= date(2025, 2, 1)
    ]
    plan42 = [
        [
            round(float(value), 4)
            for value in two._to_template_order(
                item.grid_kwh, chronological_order
            )
        ]
        + [round(float(np.sum(item.grid_kwh)), 4), round(item.plan_cost_yuan, 4)]
        for item in selected42
    ]
    plan43 = [
        [
            round(float(value), 4)
            for value in two._to_template_order(
                item.plan_grid_kwh, chronological_order
            )
        ]
        + [
            round(float(np.sum(item.plan_grid_kwh)), 4),
            round(item.plan_cost_yuan, 4),
        ]
        for item in selected43_results
    ]
    adjust43 = [
        [
            round(float(value), 4)
            for value in two._to_template_order(
                item.final_grid_kwh, chronological_order
            )
        ]
        + [
            round(float(np.sum(item.final_grid_kwh)), 4),
            round(item.settled_purchase_cost_yuan, 4),
        ]
        for item in selected43_results
    ]
    return {
        "q42": {
            "plan": plan42,
            "storage": _storage_rows(selected42),
            "emergency": _emergency_rows(
                selected42, cfg.emergency_tolerance_kwh
            ),
        },
        "q43": {
            "plan": plan43,
            "adjust": adjust43,
            "storage": _storage_rows(selected43_results, q43=True),
            "emergency": _emergency_rows(
                selected43_results, cfg.emergency_tolerance_kwh
            ),
        },
    }


def write_csv_outputs(
    results42: list[Daily42],
    summaries43: list[three.StrategySummary],
    selected43: three.StrategySummary,
    prices: PriceData,
    price_forecasts: list[PriceForecast],
    args: argparse.Namespace,
) -> None:
    with args.summary42.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file, lineterminator="\n")
        writer.writerow(
            [
                "日期",
                "计划购电量(kWh)",
                "计划购电费(元)",
                "紧急购电量(kWh)",
                "紧急购电费(元)",
                "总购电费(元)",
                "0:00储电量(kWh)",
                "24:00储电量(kWh)",
                "电价预测MAE(元/kWh)",
            ]
        )
        for index, item in enumerate(results42):
            writer.writerow(
                [
                    item.day.isoformat(),
                    f"{np.sum(item.grid_kwh):.6f}",
                    f"{item.plan_cost_yuan:.6f}",
                    f"{np.sum(item.emergency_kwh):.6f}",
                    f"{item.emergency_cost_yuan:.6f}",
                    f"{item.total_cost_yuan:.6f}",
                    f"{item.soc_start_kwh:.6f}",
                    f"{item.soc_end_kwh:.6f}",
                    f"{np.mean(np.abs(prices.price[index] - item.forecast_price)):.8f}",
                ]
            )
    three.write_daily_summary(args.summary43, selected43)
    three.write_strategy_comparison(args.comparison, summaries43)
    with args.price_evaluation.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file, lineterminator="\n")
        writer.writerow(
            [
                "日期",
                "MAE(元/kWh)",
                "RMSE(元/kWh)",
                "MAPE",
                "前日偏差权重",
                "实际均价(元/kWh)",
                "预测均价(元/kWh)",
            ]
        )
        for index, forecast in enumerate(price_forecasts):
            actual = prices.price[index]
            error = actual - forecast.value
            writer.writerow(
                [
                    prices.dates[index].isoformat(),
                    f"{np.mean(np.abs(error)):.8f}",
                    f"{np.sqrt(np.mean(error**2)):.8f}",
                    f"{np.mean(np.abs(error) / np.maximum(actual, 0.05)):.8%}",
                    f"{forecast.bias_weight:.8f}",
                    f"{np.mean(actual):.8f}",
                    f"{np.mean(forecast.value):.8f}",
                ]
            )


def _price_period_metrics(
    prices: PriceData,
    price_forecasts: list[PriceForecast],
    indices: list[int],
) -> dict[str, float]:
    actual = prices.price[indices]
    forecast = np.asarray([price_forecasts[index].value for index in indices])
    error = actual - forecast
    return {
        "电价MAE(元/kWh)": float(np.mean(np.abs(error))),
        "电价RMSE(元/kWh)": float(np.sqrt(np.mean(error**2))),
        "电价修正MAPE": float(
            np.mean(np.abs(error) / np.maximum(actual, 0.05))
        ),
    }


def _q42_period_metrics(results: list[Daily42]) -> dict[str, float]:
    plan_energy = float(sum(np.sum(item.grid_kwh) for item in results))
    emergency_energy = float(sum(np.sum(item.emergency_kwh) for item in results))
    plan_cost = float(sum(item.plan_cost_yuan for item in results))
    emergency_cost = float(sum(item.emergency_cost_yuan for item in results))
    purchased = plan_energy + emergency_energy
    return {
        "计划或最终购电量(kWh)": plan_energy,
        "紧急购电量(kWh)": emergency_energy,
        "紧急购电占比": 0.0 if purchased <= 0.0 else emergency_energy / purchased,
        "常规或调整购电费(元)": plan_cost,
        "紧急购电费(元)": emergency_cost,
        "总购电费(元)": plan_cost + emergency_cost,
    }


def write_split_evaluation(
    path: Path,
    prices: PriceData,
    price_forecasts: list[PriceForecast],
    results42: list[Daily42],
    summaries43: list[three.StrategySummary],
    selected43: three.StrategySummary,
    split: two.DatasetSplit,
) -> None:
    rows: list[dict[str, object]] = []
    for label in ("训练集", "验证集", "测试集"):
        indices = split.cycle_labels(prices.dates)[label]
        if not indices:
            raise ValueError(f"问题四的{label}为空。")
        common: dict[str, object] = {
            "数据集": label,
            "开始日期": prices.dates[indices[0]].isoformat(),
            "结束日期": prices.dates[indices[-1]].isoformat(),
            "天数": len(indices),
        }
        price_metrics = _price_period_metrics(prices, price_forecasts, indices)
        price_row = {
            **common,
            "模型或策略": "动态电价预测",
            "验证集选定策略": "",
            **price_metrics,
        }
        rows.append(price_row)

        q42_selected = [results42[index] for index in indices]
        rows.append(
            {
                **common,
                "模型或策略": "问题4-2_动态电价LP",
                "验证集选定策略": "",
                **_q42_period_metrics(q42_selected),
            }
        )
        for summary in summaries43:
            metrics43 = three.strategy_period_metrics(summary, split, label)
            rows.append(
                {
                    **common,
                    "模型或策略": f"问题4-3_{summary.name}",
                    "验证集选定策略": "是" if summary.name == selected43.name else "否",
                    "计划或最终购电量(kWh)": metrics43["最终购电量(kWh)"],
                    "紧急购电量(kWh)": metrics43["紧急购电量(kWh)"],
                    "紧急购电占比": metrics43["紧急购电占比"],
                    "常规或调整购电费(元)": metrics43["调整后购电费(元)"],
                    "紧急购电费(元)": metrics43["紧急购电费(元)"],
                    "总购电费(元)": metrics43["总购电费(元)"],
                }
            )

    headers = [
        "数据集",
        "开始日期",
        "结束日期",
        "天数",
        "模型或策略",
        "验证集选定策略",
        "电价MAE(元/kWh)",
        "电价RMSE(元/kWh)",
        "电价修正MAPE",
        "计划或最终购电量(kWh)",
        "紧急购电量(kWh)",
        "紧急购电占比",
        "常规或调整购电费(元)",
        "紧急购电费(元)",
        "总购电费(元)",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=headers, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)

    validation = three.strategy_period_metrics(selected43, split, "验证集")
    test = three.strategy_period_metrics(selected43, split, "测试集")
    print("\n问题四训练/验证/测试检验")
    print(
        f"  验证集选定策略：{selected43.name}，"
        f"验证集费用={validation['总购电费(元)']:.4f} 元"
    )
    print(f"  冻结策略后的测试集费用：{test['总购电费(元)']:.4f} 元")
    print(f"数据集检验报告：{path.resolve()}")


def write_workbooks(
    payload: dict[str, object], args: argparse.Namespace
) -> None:
    def copy_style(source, target) -> None:
        target._style = copy(source._style)
        target.number_format = source.number_format
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)

    def write_rows(sheet, rows, column_count: int, style_period: int) -> None:
        style_count = min(style_period, max(sheet.max_row - 1, 1))
        styles = [
            [copy(sheet.cell(row, column)) for column in range(1, column_count + 1)]
            for row in range(2, 2 + style_count)
        ]
        if sheet.max_row > 1:
            sheet.delete_rows(2, sheet.max_row - 1)
        for row_index, values in enumerate(rows, start=2):
            style_row = styles[(row_index - 2) % style_count]
            for column in range(1, column_count + 1):
                target = sheet.cell(row_index, column)
                copy_style(style_row[column - 1], target)
                value = values[column - 1]
                if column == 1 and isinstance(value, str):
                    try:
                        value = date.fromisoformat(value)
                    except ValueError:
                        pass
                target.value = value

    def build(
        template: Path,
        output: Path,
        section: dict[str, object],
        q43: bool,
    ) -> None:
        if not template.exists():
            raise FileNotFoundError(f"找不到结果模板：{template}")
        if template.resolve() == output.resolve():
            raise ValueError("输出路径不能覆盖官方模板。")
        output.parent.mkdir(parents=True, exist_ok=True)
        temporary = output.with_name(f".{output.stem}.building{output.suffix}")
        temporary.unlink(missing_ok=True)
        shutil.copy2(template, temporary)
        try:
            workbook = load_workbook(temporary)
            plan = workbook["计划购电量"]
            for row, values in enumerate(section["plan"], start=2):
                for column, value in enumerate(values, start=2):
                    plan.cell(row, column, value)
            if q43:
                adjust = workbook["调整购电量"]
                for row, values in enumerate(section["adjust"], start=2):
                    for column, value in enumerate(values, start=2):
                        adjust.cell(row, column, value)
            write_rows(
                workbook["充放电量"],
                section["storage"],
                6,
                24 if q43 else 18,
            )
            write_rows(workbook["紧急购电量"], section["emergency"], 3, 2)
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
            workbook.save(temporary)
            try:
                os.replace(temporary, output)
            except PermissionError as exc:
                raise PermissionError(
                    f"无法写入{output}，请关闭正在打开该文件的Excel窗口后重试。"
                ) from exc
        finally:
            temporary.unlink(missing_ok=True)

    build(args.template42, args.output42, payload["q42"], q43=False)
    build(args.template43, args.output43, payload["q43"], q43=True)


def print_summary(
    results42: list[Daily42],
    summaries43: list[three.StrategySummary],
    selected43: three.StrategySummary,
    prices: PriceData,
    price_forecasts: list[PriceForecast],
    args: argparse.Namespace,
) -> None:
    output42 = [item for item in results42 if item.day >= date(2025, 2, 1)]
    forecast_matrix = np.asarray(
        [item.value for item in price_forecasts[31:]], dtype=float
    )
    actual_matrix = prices.price[31:]
    mae = float(np.mean(np.abs(actual_matrix - forecast_matrix)))
    rmse = float(np.sqrt(np.mean((actual_matrix - forecast_matrix) ** 2)))
    total42 = float(sum(item.total_cost_yuan for item in output42))
    emergency42 = float(sum(np.sum(item.emergency_kwh) for item in output42))
    print("\n问题四动态电价预测（2025-02-01至2025-12-31）")
    print(f"  MAE={mae:.6f} 元/kWh，RMSE={rmse:.6f} 元/kWh")
    print(f"  问题4-2总费用={total42:.4f} 元，紧急购电={emergency42:.4f} kWh")
    print("\n问题4-3策略比较")
    for item in summaries43:
        print(
            f"  {item.name}：总费用={item.total_cost_yuan:.4f} 元，"
            f"紧急购电={item.emergency_energy_kwh:.4f} kWh"
        )
    print(f"  正式采用：{selected43.name}")
    print(f"\n结果文件：{args.output42.resolve()}")
    print(f"结果文件：{args.output43.resolve()}")


@dataclass(frozen=True)
class PriceSearchPoint:
    """动态电价贝叶斯优化的一个超参数候选。"""

    history_days: int
    half_life_days: float
    same_weekday_multiplier: float
    fourier_harmonics: int
    wavelet_level: int
    bias_days: int
    default_bias_weight: float
    hampel_radius: int
    hampel_threshold: float
    hampel_strength: float


def _default_price_search_point() -> PriceSearchPoint:
    cfg = Config()
    return PriceSearchPoint(
        cfg.price_history_days,
        cfg.price_half_life_days,
        cfg.price_same_weekday_multiplier,
        cfg.price_fourier_harmonics,
        cfg.price_wavelet_level,
        cfg.price_bias_days,
        cfg.default_bias_weight,
        3,
        3.0,
        0.0,
    )


def _decode_price_point(
    unit: np.ndarray, disable_hampel: bool
) -> PriceSearchPoint:
    x = np.clip(np.asarray(unit, dtype=float), 0.0, 1.0)
    return PriceSearchPoint(
        history_days=int(round(14 + 46 * x[0])),
        half_life_days=float(3.0 + 27.0 * x[1]),
        same_weekday_multiplier=float(5.0 * x[2]),
        fourier_harmonics=int(round(2 + 8 * x[3])),
        wavelet_level=int(round(1 + 3 * x[4])),
        bias_days=int(round(7 + 53 * x[5])),
        default_bias_weight=float(x[6]),
        hampel_radius=3 if disable_hampel else int(round(2 + 8 * x[7])),
        hampel_threshold=3.0 if disable_hampel else float(2.0 + 3.0 * x[8]),
        hampel_strength=0.0 if disable_hampel else float(0.25 + 0.75 * x[9]),
    )


def _encode_price_point(
    point: PriceSearchPoint, disable_hampel: bool
) -> np.ndarray:
    encoded = np.asarray(
        [
            (point.history_days - 14) / 46,
            (point.half_life_days - 3.0) / 27.0,
            point.same_weekday_multiplier / 5.0,
            (point.fourier_harmonics - 2) / 8,
            (point.wavelet_level - 1) / 3,
            (point.bias_days - 7) / 53,
            point.default_bias_weight,
            (point.hampel_radius - 2) / 8,
            (point.hampel_threshold - 2.0) / 3.0,
            max(0.0, (point.hampel_strength - 0.25) / 0.75),
        ],
        dtype=float,
    )
    return encoded[:7] if disable_hampel else encoded


def _price_point_key(point: PriceSearchPoint) -> tuple[object, ...]:
    return (
        point.history_days,
        round(point.half_life_days, 6),
        round(point.same_weekday_multiplier, 6),
        point.fourier_harmonics,
        point.wavelet_level,
        point.bias_days,
        round(point.default_bias_weight, 6),
        point.hampel_radius,
        round(point.hampel_threshold, 6),
        round(point.hampel_strength, 6),
    )


def _hampel_filter_price_history(
    matrix: np.ndarray,
    radius: int,
    threshold: float,
    strength: float,
) -> tuple[np.ndarray, int]:
    if strength <= 0.0 or matrix.shape[0] < 3:
        return matrix.copy(), 0
    filtered = matrix.copy()
    replacements = 0
    for row in range(matrix.shape[0]):
        left = max(0, row - radius)
        right = min(matrix.shape[0], row + radius + 1)
        window = matrix[left:right]
        median = np.median(window, axis=0)
        mad = np.median(np.abs(window - median), axis=0)
        robust_sigma = 1.4826 * mad
        mask = (robust_sigma > 1.0e-12) & (
            np.abs(matrix[row] - median) > threshold * robust_sigma
        )
        if np.any(mask):
            filtered[row, mask] = (
                (1.0 - strength) * matrix[row, mask] + strength * median[mask]
            )
            replacements += int(np.count_nonzero(mask))
    return filtered, replacements


def _search_price_history_profile(
    prices: PriceData,
    cutoff: int,
    target_date: date,
    baseline_price: np.ndarray,
    point: PriceSearchPoint,
    training_indices: set[int],
) -> tuple[np.ndarray, int]:
    first = max(0, cutoff - point.history_days)
    indices = np.asarray(
        [index for index in range(first, cutoff) if index in training_indices],
        dtype=int,
    )
    if indices.size == 0:
        return baseline_price.copy(), 0
    matrix, replacements = _hampel_filter_price_history(
        prices.price[indices],
        point.hampel_radius,
        point.hampel_threshold,
        point.hampel_strength,
    )
    ages = cutoff - 1 - indices
    weights = np.exp(-math.log(2.0) * ages / point.half_life_days)
    same_weekday = np.asarray(
        [prices.dates[index].weekday() == target_date.weekday() for index in indices],
        dtype=float,
    )
    weights *= 1.0 + point.same_weekday_multiplier * same_weekday
    return np.average(matrix, axis=0, weights=weights), replacements


def _search_price_bases(
    prices: PriceData,
    baseline_price: np.ndarray,
    point: PriceSearchPoint,
    training_indices: set[int],
) -> tuple[np.ndarray, int]:
    hours = np.arange(144, dtype=float) / 6.0
    bases: list[np.ndarray] = []
    replacements = 0
    for cutoff, day in enumerate(prices.dates):
        profile, count = _search_price_history_profile(
            prices, cutoff, day, baseline_price, point, training_indices
        )
        periodic = fit_price_fourier(hours, profile, point.fourier_harmonics)
        wavelet = two.wavelet_correct_residual(profile, periodic, point.wavelet_level)
        bases.append(periodic + wavelet)
        replacements += count
    return np.asarray(bases), replacements


def _search_bias_weight(
    prices: PriceData,
    bases: np.ndarray,
    cutoff: int,
    point: PriceSearchPoint,
    training_indices: set[int],
) -> float:
    pair_end = cutoff - 1
    pair_start = max(1, pair_end - point.bias_days + 1)
    previous: list[np.ndarray] = []
    current: list[np.ndarray] = []
    for index in range(pair_start, pair_end + 1):
        if index in training_indices and index - 1 in training_indices:
            previous.append(prices.price[index - 1] - bases[index - 1])
            current.append(prices.price[index] - bases[index])
    if len(previous) < Config().minimum_bias_pairs:
        return point.default_bias_weight
    x = np.concatenate(previous)
    y = np.concatenate(current)
    denominator = float(np.dot(x, x))
    if denominator <= 1.0e-12:
        return point.default_bias_weight
    return float(np.clip(np.dot(x, y) / denominator, 0.0, 1.0))


def _search_price_forecasts(
    prices: PriceData,
    baseline_price: np.ndarray,
    point: PriceSearchPoint,
    training_indices: set[int],
) -> tuple[np.ndarray, int]:
    hours = np.arange(144, dtype=float) / 6.0
    bases, replacements = _search_price_bases(
        prices, baseline_price, point, training_indices
    )
    forecasts: list[np.ndarray] = []
    for cutoff, day in enumerate(prices.dates):
        profile, _ = _search_price_history_profile(
            prices, cutoff, day, baseline_price, point, training_indices
        )
        periodic = fit_price_fourier(hours, profile, point.fourier_harmonics)
        wavelet = two.wavelet_correct_residual(profile, periodic, point.wavelet_level)
        latest = cutoff - 1
        if latest >= 0 and latest in training_indices:
            bias = prices.price[latest] - bases[latest]
            weight = _search_bias_weight(
                prices, bases, cutoff, point, training_indices
            )
        else:
            bias = np.zeros(144, dtype=float)
            weight = 0.0
        forecasts.append(np.maximum(0.0, periodic + wavelet + weight * bias))
    return np.asarray(forecasts), replacements


def _bayes_price_metrics(
    actual: np.ndarray, forecast: np.ndarray
) -> dict[str, float]:
    error = actual - forecast
    denominator = np.maximum(np.abs(actual), 0.05)
    residual = float(np.sum(error**2))
    total = float(np.sum((actual - np.mean(actual)) ** 2))
    return {
        "mae": float(np.mean(np.abs(error))),
        "rmse": float(np.sqrt(np.mean(error**2))),
        "mape": float(np.mean(np.abs(error) / denominator)),
        "r2": float("nan") if total <= 0.0 else 1.0 - residual / total,
    }


def run_price_bayesian_optimization(args: argparse.Namespace) -> None:
    """严格前向电价寻优；可通过开关执行Hampel消融实验。"""
    if args.bayes_trials < 2 or not 1 <= args.bayes_initial < args.bayes_trials:
        raise ValueError("应满足 bayes-trials >= 2 且 1 <= bayes-initial < bayes-trials。")
    report_path = args.bayes_report
    csv_path = args.bayes_csv
    if args.bayes_disable_hampel:
        if report_path == DEFAULT_PRICE_BAYES_REPORT:
            report_path = DEFAULT_PRICE_BAYES_NO_HAMPEL_REPORT
        if csv_path == DEFAULT_PRICE_BAYES_CSV:
            csv_path = DEFAULT_PRICE_BAYES_NO_HAMPEL_CSV

    baseline, order = two.read_baseline(args.baseline)
    prices = read_dynamic_prices(args.price, order)
    groups = two.DatasetSplit().cycle_labels(prices.dates)
    training_indices = set(groups["训练集"])
    labels = {
        index: label for label, indices in groups.items() for index in indices
    }
    rng = np.random.default_rng(args.bayes_seed)
    base_point = _default_price_search_point()
    candidates = [base_point]
    base_prediction, base_replacements = _search_price_forecasts(
        prices, baseline.price, base_point, training_indices
    )
    predictions = [base_prediction]
    replacement_counts = [base_replacements]
    seen = {_price_point_key(base_point)}
    validation_seen: list[int] = []
    current_index = 0
    forward_prediction: list[np.ndarray] = []
    selection_history: list[dict[str, object]] = []
    proposals: list[dict[str, object]] = []
    proposal_count = 1
    no_improvement = 0
    search_stopped = False

    print("严格前向电价贝叶斯优化：测试集不参与参数搜索。")
    for index, day in enumerate(prices.dates):
        forward_prediction.append(predictions[current_index][index])
        if labels[index] != "验证集":
            continue
        validation_seen.append(index)
        actual_seen = prices.price[validation_seen]
        scores = np.asarray(
            [
                _bayes_price_metrics(actual_seen, prediction[validation_seen])["mape"]
                for prediction in predictions
            ]
        )
        incumbent = float(np.min(scores))
        if proposal_count < args.bayes_trials and not search_stopped:
            dimensions = 7 if args.bayes_disable_hampel else 10
            if len(candidates) < args.bayes_initial:
                unit = rng.random(dimensions)
                method = "随机初始化"
            else:
                pool = rng.random((args.bayes_candidate_pool, dimensions))
                acquisition = two._expected_improvement(
                    np.asarray(
                        [
                            _encode_price_point(point, args.bayes_disable_hampel)
                            for point in candidates
                        ]
                    ),
                    scores,
                    pool,
                    0.42,
                )
                unit = pool[int(np.argmax(acquisition))]
                method = "贝叶斯期望改进"
            point = _decode_price_point(unit, args.bayes_disable_hampel)
            attempts = 0
            while _price_point_key(point) in seen and attempts < 100:
                point = _decode_price_point(
                    rng.random(dimensions), args.bayes_disable_hampel
                )
                attempts += 1
            seen.add(_price_point_key(point))
            proposal_count += 1
            prediction, replacement_count = _search_price_forecasts(
                prices, baseline.price, point, training_indices
            )
            candidate_score = _bayes_price_metrics(
                actual_seen, prediction[validation_seen]
            )["mape"]
            improved = candidate_score < incumbent - args.bayes_min_delta
            candidates.append(point)
            predictions.append(prediction)
            replacement_counts.append(replacement_count)
            no_improvement = 0 if improved else no_improvement + int(
                proposal_count >= args.bayes_initial
            )
            proposals.append(
                {
                    "proposal": proposal_count,
                    "validation_date": day.isoformat(),
                    "validation_days_available": len(validation_seen),
                    "method": method,
                    "validation_mape": candidate_score,
                    "improved": improved,
                    "hampel_replacements_in_all_causal_windows": replacement_count,
                    "parameters": asdict(point),
                }
            )
            print(f"[{day}] 累计验证MAPE={candidate_score:.4%}")
            if len(candidates) >= args.bayes_initial and no_improvement >= args.bayes_patience:
                search_stopped = True
                print(f"连续{args.bayes_patience}个候选无显著改善，停止寻优。")

        scores = np.asarray(
            [
                _bayes_price_metrics(actual_seen, prediction[validation_seen])["mape"]
                for prediction in predictions
            ]
        )
        selected = int(np.argmin(scores))
        if len(validation_seen) >= args.bayes_initial:
            current_index = selected
        selection_history.append(
            {
                "date": day.isoformat(),
                "validation_days_available": len(validation_seen),
                "selected_candidate": current_index,
                "selected_validation_mape": float(scores[current_index]),
                "selected_parameters": asdict(candidates[current_index]),
            }
        )

    forward = np.asarray(forward_prediction)
    baseline_metrics = {
        label: _bayes_price_metrics(prices.price[indices], base_prediction[indices])
        for label, indices in groups.items()
    }
    forward_metrics = {
        label: _bayes_price_metrics(prices.price[indices], forward[indices])
        for label, indices in groups.items()
    }
    report = {
        "data_split": {
            "rule": "训练、训练、训练、验证、训练、测试",
            "counts": {label: len(indices) for label, indices in groups.items()},
            "strict_training_freeze": True,
            "test_data_used_in_optimization": False,
        },
        "optimization": {
            "method": "严格前向高斯过程贝叶斯优化 + 早停",
            "hampel_enabled": not args.bayes_disable_hampel,
            "objective": "截至当前日期已开放验证集的累计MAPE",
            "maximum_proposals": args.bayes_trials,
            "completed_proposals": proposal_count,
            "patience": args.bayes_patience,
            "minimum_absolute_mape_improvement": args.bayes_min_delta,
            "search_early_stopped": search_stopped,
            "seed": args.bayes_seed,
        },
        "baseline_formal_parameters_metrics": baseline_metrics,
        "strict_forward_experiment_metrics": forward_metrics,
        "year_end_selected_parameters": asdict(candidates[current_index]),
        "year_end_hampel_replacements_in_all_causal_windows": replacement_counts[current_index],
        "proposals": proposals,
        "selection_history": selection_history,
    }
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    with csv_path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file, lineterminator="\n")
        writer.writerow(
            [
                "日期",
                "数据集",
                "实际均价",
                "正式模型预测均价",
                "实验模型预测均价",
                "正式模型MAPE",
                "实验模型MAPE",
            ]
        )
        for index, day in enumerate(prices.dates):
            actual = prices.price[index]
            base = base_prediction[index]
            experimental = forward[index]
            denominator = np.maximum(np.abs(actual), 0.05)
            writer.writerow(
                [
                    day.isoformat(),
                    labels[index],
                    f"{np.mean(actual):.8f}",
                    f"{np.mean(base):.8f}",
                    f"{np.mean(experimental):.8f}",
                    f"{np.mean(np.abs(actual - base) / denominator):.8f}",
                    f"{np.mean(np.abs(actual - experimental) / denominator):.8f}",
                ]
            )
    for label in ("训练集", "验证集", "测试集"):
        values = forward_metrics[label]
        print(
            f"{label}：MAE={values['mae']:.6f}，"
            f"RMSE={values['rmse']:.6f}，MAPE={values['mape']:.4%}"
        )
    print(f"实验报告：{report_path.resolve()}")
    print(f"逐日评价：{csv_path.resolve()}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="C题问题四傅里叶—小波动态电价预测与LP/MPC购电策略"
    )
    parser.add_argument("--baseline", type=Path, default=two.DEFAULT_BASELINE)
    parser.add_argument("--load", type=Path, default=two.DEFAULT_LOAD)
    parser.add_argument("--pv", type=Path, default=two.DEFAULT_PV)
    parser.add_argument("--forecast", type=Path, default=DEFAULT_FORECAST)
    parser.add_argument("--price", type=Path, default=DEFAULT_PRICE)
    parser.add_argument("--template42", type=Path, default=DEFAULT_TEMPLATE_42)
    parser.add_argument("--template43", type=Path, default=DEFAULT_TEMPLATE_43)
    parser.add_argument("--output42", type=Path, default=DEFAULT_OUTPUT_42)
    parser.add_argument("--output43", type=Path, default=DEFAULT_OUTPUT_43)
    parser.add_argument("--summary42", type=Path, default=DEFAULT_SUMMARY_42)
    parser.add_argument("--summary43", type=Path, default=DEFAULT_SUMMARY_43)
    parser.add_argument("--comparison", type=Path, default=DEFAULT_COMPARISON)
    parser.add_argument(
        "--price-evaluation", type=Path, default=DEFAULT_PRICE_EVALUATION
    )
    parser.add_argument("--split-report", type=Path, default=DEFAULT_SPLIT_REPORT)
    parser.add_argument("--train-end", type=date.fromisoformat, default=date(2025, 7, 31))
    parser.add_argument(
        "--validation-end", type=date.fromisoformat, default=date(2025, 9, 30)
    )
    parser.add_argument(
        "--skip-workbooks",
        action="store_true",
        help="只生成CSV与训练/验证/测试检验报告，不重写Excel结果表。",
    )
    parser.add_argument(
        "--bayes-optimize",
        action="store_true",
        help="执行严格前向电价贝叶斯寻优，不运行4-2和4-3正式调度。",
    )
    parser.add_argument("--bayes-trials", type=int, default=36)
    parser.add_argument("--bayes-initial", type=int, default=10)
    parser.add_argument("--bayes-patience", type=int, default=10)
    parser.add_argument("--bayes-min-delta", type=float, default=0.0005)
    parser.add_argument("--bayes-candidate-pool", type=int, default=2500)
    parser.add_argument("--bayes-seed", type=int, default=2026)
    parser.add_argument("--bayes-report", type=Path, default=DEFAULT_PRICE_BAYES_REPORT)
    parser.add_argument("--bayes-csv", type=Path, default=DEFAULT_PRICE_BAYES_CSV)
    parser.add_argument(
        "--bayes-disable-hampel",
        action="store_true",
        help="执行关闭Hampel滤波的贝叶斯消融实验。",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if args.bayes_optimize:
        run_price_bayesian_optimization(args)
        return
    cfg = Config()
    baseline, chronological_order = two.read_baseline(args.baseline)
    historical = two.read_historical_data(args.load, args.pv, chronological_order)
    prices = read_dynamic_prices(args.price, chronological_order)
    if prices.dates != historical.dates:
        raise ValueError("附件2与附件4的日期不一致。")
    forecasts = three.read_forecasts(args.forecast)
    split = two.DatasetSplit(args.train_end, args.validation_end)
    groups = split.cycle_labels(historical.dates)
    training_indices = set(groups["训练集"])

    print("正在构造傅里叶—小波—前日偏差动态电价预测……")
    current_prices, following_prices = prepare_daily_price_forecasts(
        prices,
        baseline.price,
        cfg,
        training_indices=training_indices,
    )
    print("正在计算问题4-2……")
    results42 = simulate_42(
        baseline,
        historical,
        prices,
        current_prices,
        following_prices,
        cfg,
    )
    validate_42(results42, cfg)

    print("正在构造问题4-3负荷、光伏决策输入……")
    decision_inputs = three.prepare_decision_inputs(
        baseline,
        historical,
        forecasts,
        cfg,
        training_indices=training_indices,
    )
    strategy_specs = [
        ("S0_仅0点", (0,)),
        ("S1_增加6点", (0, 6)),
        ("S2_增加12点", (0, 6, 12)),
        ("S3_完整MPC", (0, 6, 12, 18)),
    ]
    summaries43: list[three.StrategySummary] = []
    for name, release_hours in strategy_specs:
        print(f"正在计算问题4-3策略{name}……")
        summary = simulate_43_strategy(
            name,
            release_hours,
            historical,
            prices,
            decision_inputs,
            current_prices,
            following_prices,
            cfg,
        )
        three.validate_strategy(summary, cfg)
        summaries43.append(summary)
    selected43 = three.select_strategy_on_validation(summaries43, split)

    for path in (
        args.summary42,
        args.summary43,
        args.comparison,
        args.price_evaluation,
        args.split_report,
    ):
        path.parent.mkdir(parents=True, exist_ok=True)
    write_csv_outputs(
        results42,
        summaries43,
        selected43,
        prices,
        current_prices,
        args,
    )
    write_split_evaluation(
        args.split_report,
        prices,
        current_prices,
        results42,
        summaries43,
        selected43,
        split,
    )
    if not args.skip_workbooks:
        payload = make_workbook_payload(
            results42, selected43, chronological_order, cfg
        )
        write_workbooks(payload, args)
    print_summary(
        results42,
        summaries43,
        selected43,
        prices,
        current_prices,
        args,
    )


if __name__ == "__main__":
    main()
