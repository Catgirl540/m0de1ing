"""2026 CUMCM C题问题一：线性规划与傅里叶—小波—正弦拟合对照。

主模型直接使用附件1的原始10分钟数据求最优购电策略。
改进模型先以二阶傅里叶级数提取负载的日周期趋势，再用三级 Haar
小波软阈值重构傅里叶残差中的局部变化；光伏仍采用截断幂正弦拟合。
随后用同一线性规划求解，并同时报告纯傅里叶模型和小波改进模型。

默认将傅里叶—小波负荷和截断正弦光伏对应的策略写入 result1.xlsx；
仍可通过 --result-scenario actual 输出原始数据理论最优解。

默认输入：C题/附件/csv/附件1.csv
默认模板：C题/附件/附件5/result1.xlsx
默认输出：C题/result1.xlsx
默认对照：C题/问题一拟合对比.csv

依赖：numpy、scipy、openpyxl（Haar 小波由本文件实现，无需 PyWavelets）
运行：python one.py
"""

from __future__ import annotations

import argparse
import csv
import math
import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
DEFAULT_INPUT = ROOT / "C题" / "附件" / "csv" / "附件1.csv"
DEFAULT_TEMPLATE = ROOT / "C题" / "附件" / "附件5" / "result1.xlsx"
DEFAULT_OUTPUT = ROOT / "C题" / "result1.xlsx"
DEFAULT_COMPARISON = ROOT / "C题" / "问题一拟合对比.csv"


@dataclass(frozen=True)
class Config:
    dt_hours: float = 1.0 / 6.0
    capacity_kwh: float = 12000.0
    soc_min_kwh: float = 1200.0
    soc_max_kwh: float = 10800.0
    initial_soc_kwh: float = 6000.0
    max_charge_power_kw: float = 5000.0
    max_discharge_power_kw: float = 5000.0
    charge_efficiency: float = 0.90
    discharge_efficiency: float = 0.90
    # 同购电费方案中优先选择较少充放电的方案，不改变主目标。
    throughput_tie_breaker: float = 1.0e-7

    @property
    def max_charge_kwh(self) -> float:
        return self.max_charge_power_kw * self.dt_hours

    @property
    def max_discharge_kwh(self) -> float:
        return self.max_discharge_power_kw * self.dt_hours


@dataclass
class InputData:
    times: list[str]
    price: np.ndarray
    load_kw: np.ndarray
    pv_kw: np.ndarray


@dataclass
class FitResult:
    load_fourier_kw: np.ndarray
    load_kw: np.ndarray
    pv_kw: np.ndarray
    load_fourier_mae: float
    load_fourier_rmse: float
    load_fourier_r2: float
    load_mae: float
    load_rmse: float
    load_r2: float
    pv_mae: float
    pv_rmse: float
    pv_r2: float
    pv_amplitude_kw: float
    sunrise_hour: float
    sunset_hour: float
    pv_shape: float
    wavelet_name: str
    wavelet_level: int


@dataclass
class Solution:
    grid_kwh: np.ndarray
    charge_kwh: np.ndarray
    discharge_kwh: np.ndarray
    curtailment_kwh: np.ndarray
    soc_kwh: np.ndarray
    total_cost_yuan: float


def _as_float(value: object, field: str, row: int) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"附件1第{row}行的{field}不是数值：{value!r}") from exc
    if not math.isfinite(number):
        raise ValueError(f"附件1第{row}行的{field}不是有限数值：{value!r}")
    return number


def read_input(path: Path) -> InputData:
    """从UTF-8 CSV读取时间、电价、负载与光伏功率。"""
    if not path.exists():
        raise FileNotFoundError(f"找不到附件1 CSV：{path}")
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        rows = list(csv.reader(file))
    if len(rows) != 145 or len(rows[0]) < 4:
        raise ValueError("附件1应包含1行表头、144行数据和至少4列。")

    times, price, load, pv = [], [], [], []
    for row_number, row in enumerate(rows[1:], start=2):
        if len(row) < 4:
            raise ValueError(f"附件1第{row_number}行少于4列。")
        times.append(row[0].strip())
        price.append(_as_float(row[1], "电价", row_number))
        load.append(_as_float(row[2], "小区负载", row_number))
        pv.append(_as_float(row[3], "光伏预测功率", row_number))
    arrays = [np.asarray(values, dtype=float) for values in (price, load, pv)]
    if any(np.any(values < 0) for values in arrays):
        raise ValueError("电价、负载和光伏功率均不应为负数。")
    return InputData(times, arrays[0], arrays[1], arrays[2])


def _start_minute(interval_label: object) -> int:
    label = str(interval_label).strip()
    if "-" not in label:
        raise ValueError(f"无法解析时间段：{label!r}")
    start = label.split("-", 1)[0].replace("+1", "")
    hour_text, minute_text = start.split(":", 1)
    return (int(hour_text) % 24) * 60 + int(minute_text)


def template_chronological_order(template: Path) -> tuple[list[int], list[str]]:
    """把模板的0:10起始行序转换为0:00至23:50的自然日顺序。"""
    if not template.exists():
        raise FileNotFoundError(f"找不到result1模板：{template}")
    workbook = load_workbook(template, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    labels = [sheet.cell(row, 1).value for row in range(2, sheet.max_row + 1)]
    workbook.close()
    if len(labels) != 144 or any(label is None for label in labels):
        raise ValueError("result1模板应包含144个完整时间段。")
    minutes = [_start_minute(label) for label in labels]
    if sorted(minutes) != list(range(0, 1440, 10)):
        raise ValueError("result1模板未完整覆盖144个10分钟时段。")
    order = sorted(range(144), key=minutes.__getitem__)
    return order, [str(label) for label in labels]


def _metrics(actual: np.ndarray, fitted: np.ndarray) -> tuple[float, float, float]:
    residual = actual - fitted
    mae = float(np.mean(np.abs(residual)))
    rmse = float(np.sqrt(np.mean(residual**2)))
    total = float(np.sum((actual - np.mean(actual)) ** 2))
    r2 = 1.0 - float(np.sum(residual**2)) / total if total > 0 else 1.0
    return mae, rmse, r2


def fit_load_fourier(hours: np.ndarray, load_kw: np.ndarray) -> np.ndarray:
    """使用24小时、12小时两个周期项拟合早晚负荷变化。"""
    omega = 2.0 * np.pi / 24.0
    design = np.column_stack(
        [
            np.ones_like(hours),
            np.cos(omega * hours),
            np.sin(omega * hours),
            np.cos(2.0 * omega * hours),
            np.sin(2.0 * omega * hours),
        ]
    )
    coefficients, *_ = np.linalg.lstsq(design, load_kw, rcond=None)
    return np.maximum(0.0, design @ coefficients)


def _haar_dwt(signal: np.ndarray, level: int) -> tuple[np.ndarray, list[np.ndarray]]:
    """正交 Haar 离散小波分解；detail 按第1层到第 level 层保存。"""
    approximation = np.asarray(signal, dtype=float).copy()
    details: list[np.ndarray] = []
    scale = math.sqrt(2.0)
    for _ in range(level):
        if approximation.size % 2:
            raise ValueError("Haar 小波每层输入长度必须为偶数。")
        even = approximation[0::2]
        odd = approximation[1::2]
        details.append((even - odd) / scale)
        approximation = (even + odd) / scale
    return approximation, details


def _haar_idwt(approximation: np.ndarray, details: list[np.ndarray]) -> np.ndarray:
    """由 Haar 近似系数和细节系数重构信号。"""
    reconstructed = np.asarray(approximation, dtype=float).copy()
    scale = math.sqrt(2.0)
    for detail in reversed(details):
        if reconstructed.size != detail.size:
            raise ValueError("Haar 小波重构时系数长度不匹配。")
        restored = np.empty(reconstructed.size * 2, dtype=float)
        restored[0::2] = (reconstructed + detail) / scale
        restored[1::2] = (reconstructed - detail) / scale
        reconstructed = restored
    return reconstructed


def wavelet_correct_fourier_residual(
    load_kw: np.ndarray, fourier_kw: np.ndarray, level: int = 3
) -> np.ndarray:
    """对傅里叶残差做 Haar 小波软阈值重构，保留局部结构并抑制噪声。"""
    residual = np.asarray(load_kw, dtype=float) - np.asarray(fourier_kw, dtype=float)
    if residual.size % (2**level):
        raise ValueError(f"样本数必须能被 2^{level} 整除。")
    approximation, details = _haar_dwt(residual, level)
    thresholded: list[np.ndarray] = []
    n = residual.size
    for detail in details:
        # MAD 是对高斯噪声标准差的稳健估计；通用阈值避免把尖锐噪声当趋势。
        sigma = float(np.median(np.abs(detail - np.median(detail))) / 0.6745)
        threshold = sigma * math.sqrt(2.0 * math.log(n))
        thresholded.append(
            np.sign(detail) * np.maximum(np.abs(detail) - threshold, 0.0)
        )
    corrected_residual = _haar_idwt(approximation, thresholded)
    return np.maximum(0.0, fourier_kw + corrected_residual)


def fit_pv_truncated_sine(
    hours: np.ndarray, pv_kw: np.ndarray
) -> tuple[np.ndarray, float, float, float, float]:
    """拟合 A*sin(pi*(t-tr)/(ts-tr))^gamma，并在日照区间外置零。"""
    peak = float(np.max(pv_kw))
    if peak <= 0:
        return np.zeros_like(pv_kw), 0.0, 0.0, 0.0, 1.0
    positive = np.flatnonzero(pv_kw > max(1.0e-6, peak * 1.0e-5))
    first_hour = float(hours[positive[0]])
    last_hour = float(hours[positive[-1]])
    sunrise_grid = np.arange(max(0.0, first_hour - 1.5), first_hour + 0.51, 1 / 12)
    sunset_grid = np.arange(last_hour - 0.5, min(24.0, last_hour + 1.51) + 1e-9, 1 / 12)
    shape_grid = np.linspace(0.5, 3.0, 101)

    best_sse = math.inf
    best = (np.zeros_like(pv_kw), 0.0, first_hour, last_hour, 1.0)
    for sunrise in sunrise_grid:
        for sunset in sunset_grid:
            if sunset - sunrise < 6.0:
                continue
            phase = np.pi * (hours - sunrise) / (sunset - sunrise)
            sine = np.zeros_like(hours)
            daylight = (hours >= sunrise) & (hours <= sunset)
            sine[daylight] = np.maximum(0.0, np.sin(phase[daylight]))
            for shape in shape_grid:
                basis = sine**shape
                denominator = float(np.dot(basis, basis))
                if denominator <= 0:
                    continue
                amplitude = max(0.0, float(np.dot(basis, pv_kw) / denominator))
                fitted = amplitude * basis
                sse = float(np.sum((pv_kw - fitted) ** 2))
                if sse < best_sse:
                    best_sse = sse
                    best = (fitted, amplitude, float(sunrise), float(sunset), float(shape))
    return best


def fit_profiles(data: InputData, order: Iterable[int], cfg: Config) -> FitResult:
    order_array = np.asarray(list(order), dtype=int)
    hours = np.arange(144, dtype=float) * cfg.dt_hours
    load_actual = data.load_kw[order_array]
    pv_actual = data.pv_kw[order_array]
    load_fourier = fit_load_fourier(hours, load_actual)
    load_fitted = wavelet_correct_fourier_residual(
        load_actual, load_fourier, level=3
    )
    pv_fitted, amplitude, sunrise, sunset, shape = fit_pv_truncated_sine(
        hours, pv_actual
    )
    load_fourier_mae, load_fourier_rmse, load_fourier_r2 = _metrics(
        load_actual, load_fourier
    )
    load_mae, load_rmse, load_r2 = _metrics(load_actual, load_fitted)
    pv_mae, pv_rmse, pv_r2 = _metrics(pv_actual, pv_fitted)
    return FitResult(
        load_fourier,
        load_fitted,
        pv_fitted,
        load_fourier_mae,
        load_fourier_rmse,
        load_fourier_r2,
        load_mae,
        load_rmse,
        load_r2,
        pv_mae,
        pv_rmse,
        pv_r2,
        amplitude,
        sunrise,
        sunset,
        shape,
        "Haar",
        3,
    )


def solve_linear_program(
    price: np.ndarray, load_kw: np.ndarray, pv_kw: np.ndarray, cfg: Config
) -> Solution:
    """求解购电—充放电—弃光—SOC线性规划。"""
    try:
        from scipy.optimize import linprog
    except ModuleNotFoundError as exc:
        raise ModuleNotFoundError(
            "缺少scipy，请先执行：python -m pip install scipy"
        ) from exc

    n = len(price)
    if n != 144 or len(load_kw) != n or len(pv_kw) != n:
        raise ValueError("价格、负载和光伏数组都必须包含144个时段。")
    load_energy = load_kw * cfg.dt_hours
    pv_energy = pv_kw * cfg.dt_hours

    g0, c0, d0, w0, e0 = 0, n, 2 * n, 3 * n, 4 * n
    count = 5 * n + 1
    objective = np.zeros(count)
    objective[g0 : g0 + n] = price
    objective[c0 : c0 + n] = cfg.throughput_tie_breaker
    objective[d0 : d0 + n] = cfg.throughput_tie_breaker

    a_eq = np.zeros((2 * n, count), dtype=float)
    b_eq = np.zeros(2 * n, dtype=float)
    for t in range(n):
        # g + PV + d = load + c + w
        a_eq[t, g0 + t] = 1.0
        a_eq[t, c0 + t] = -1.0
        a_eq[t, d0 + t] = 1.0
        a_eq[t, w0 + t] = -1.0
        b_eq[t] = load_energy[t] - pv_energy[t]
        # e[t+1] = e[t] + eta_c*c - d/eta_d
        row = n + t
        a_eq[row, c0 + t] = -cfg.charge_efficiency
        a_eq[row, d0 + t] = 1.0 / cfg.discharge_efficiency
        a_eq[row, e0 + t] = -1.0
        a_eq[row, e0 + t + 1] = 1.0

    bounds: list[tuple[float | None, float | None]] = []
    bounds.extend([(0.0, None)] * n)
    bounds.extend([(0.0, cfg.max_charge_kwh)] * n)
    bounds.extend([(0.0, cfg.max_discharge_kwh)] * n)
    bounds.extend([(0.0, None)] * n)
    bounds.extend([(cfg.soc_min_kwh, cfg.soc_max_kwh)] * (n + 1))
    bounds[e0] = (cfg.initial_soc_kwh, cfg.initial_soc_kwh)
    bounds[e0 + n] = (cfg.initial_soc_kwh, cfg.initial_soc_kwh)

    result = linprog(
        objective,
        A_eq=a_eq,
        b_eq=b_eq,
        bounds=bounds,
        method="highs",
        options={"presolve": True},
    )
    if not result.success:
        raise RuntimeError(f"线性规划失败：status={result.status}, {result.message}")
    x = np.asarray(result.x, dtype=float)
    solution = Solution(
        x[g0 : g0 + n],
        x[c0 : c0 + n],
        x[d0 : d0 + n],
        x[w0 : w0 + n],
        x[e0 : e0 + n + 1],
        float(np.dot(price, x[g0 : g0 + n])),
    )
    validate_solution(solution, load_energy, pv_energy, price, cfg)
    return solution


def validate_solution(
    solution: Solution,
    load_energy: np.ndarray,
    pv_energy: np.ndarray,
    price: np.ndarray,
    cfg: Config,
    tolerance: float = 1.0e-5,
) -> None:
    balance = (
        solution.grid_kwh
        + pv_energy
        + solution.discharge_kwh
        - load_energy
        - solution.charge_kwh
        - solution.curtailment_kwh
    )
    next_soc = (
        solution.soc_kwh[:-1]
        + cfg.charge_efficiency * solution.charge_kwh
        - solution.discharge_kwh / cfg.discharge_efficiency
    )
    checks = {
        "电量平衡": float(np.max(np.abs(balance))),
        "SOC递推": float(np.max(np.abs(solution.soc_kwh[1:] - next_soc))),
        "费用复算": abs(solution.total_cost_yuan - float(np.dot(price, solution.grid_kwh))),
        "初始SOC": abs(solution.soc_kwh[0] - cfg.initial_soc_kwh),
        "终止SOC": abs(solution.soc_kwh[-1] - cfg.initial_soc_kwh),
    }
    failed = {name: value for name, value in checks.items() if value > tolerance}
    if failed:
        raise RuntimeError(f"求解结果校验失败：{failed}")
    if np.min(solution.soc_kwh) < cfg.soc_min_kwh - tolerance:
        raise RuntimeError("SOC低于下限。")
    if np.max(solution.soc_kwh) > cfg.soc_max_kwh + tolerance:
        raise RuntimeError("SOC高于上限。")
    if np.max(np.minimum(solution.charge_kwh, solution.discharge_kwh)) > 1.0e-4:
        raise RuntimeError("出现同时充放电，应改用MILP加入互斥变量。")


def evaluate_fitted_dispatch(
    fitted: Solution,
    actual_price: np.ndarray,
    actual_load_kw: np.ndarray,
    actual_pv_kw: np.ndarray,
    cfg: Config,
) -> tuple[np.ndarray, np.ndarray, float]:
    """保持拟合模型的储能动作不变，在真实数据下复算购电和弃光。"""
    net_need = (
        actual_load_kw * cfg.dt_hours
        + fitted.charge_kwh
        - actual_pv_kw * cfg.dt_hours
        - fitted.discharge_kwh
    )
    grid = np.maximum(0.0, net_need)
    curtailment = np.maximum(0.0, -net_need)
    return grid, curtailment, float(np.dot(actual_price, grid))


def _to_template_order(values: np.ndarray, order: list[int]) -> np.ndarray:
    mapped = np.empty_like(values)
    mapped[np.asarray(order, dtype=int)] = values
    return mapped


def write_result(template: Path, output: Path, order: list[int], solution: Solution) -> None:
    """保留官方模板格式，写入所选场景的最优解。"""
    output.parent.mkdir(parents=True, exist_ok=True)
    if template.resolve() == output.resolve():
        raise ValueError("输出路径不能覆盖官方模板。")
    shutil.copy2(template, output)
    workbook = load_workbook(output)
    plan_sheet = workbook["计划购电量"]
    storage_sheet = workbook["充放电量"]
    for row, value in enumerate(_to_template_order(solution.grid_kwh, order), start=2):
        plan_sheet.cell(row, 2, round(float(value), 4))
    for block in range(6):
        start, stop = block * 24, (block + 1) * 24
        charge = round(float(np.sum(solution.charge_kwh[start:stop])), 4)
        discharge = round(float(np.sum(solution.discharge_kwh[start:stop])), 4)
        storage_sheet.cell(2 + block, 2, charge)
        storage_sheet.cell(2 + block, 3, discharge)
    storage_sheet["E2"] = round(float(solution.soc_kwh[0]), 4)
    storage_sheet["E3"] = round(float(solution.soc_kwh[-1]), 4)
    workbook.save(output)


def write_comparison(
    output: Path,
    price: np.ndarray,
    load_actual: np.ndarray,
    pv_actual: np.ndarray,
    fit: FitResult,
    actual_solution: Solution,
    fourier_solution: Solution,
    wavelet_solution: Solution,
    fourier_grid_actual: np.ndarray,
    wavelet_grid_actual: np.ndarray,
) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    headers = [
        "时间",
        "电价(元/kWh)",
        "实际负载(kW)",
        "傅里叶拟合负载(kW)",
        "傅里叶-小波拟合负载(kW)",
        "实际光伏功率(kW)",
        "截断正弦拟合光伏功率(kW)",
        "原始数据最优购电量(kWh)",
        "傅里叶场景计划购电量(kWh)",
        "小波改进场景计划购电量(kWh)",
        "傅里叶策略在真实数据下购电量(kWh)",
        "小波改进策略在真实数据下购电量(kWh)",
        "原始策略充电量(kWh)",
        "原始策略放电量(kWh)",
        "原始策略SOC期末值(kWh)",
    ]
    with output.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file, lineterminator="\n")
        writer.writerow(headers)
        for t in range(144):
            writer.writerow(
                [
                    f"{t // 6:02d}:{(t % 6) * 10:02d}",
                    f"{price[t]:.6f}",
                    f"{load_actual[t]:.6f}",
                    f"{fit.load_fourier_kw[t]:.6f}",
                    f"{fit.load_kw[t]:.6f}",
                    f"{pv_actual[t]:.6f}",
                    f"{fit.pv_kw[t]:.6f}",
                    f"{actual_solution.grid_kwh[t]:.6f}",
                    f"{fourier_solution.grid_kwh[t]:.6f}",
                    f"{wavelet_solution.grid_kwh[t]:.6f}",
                    f"{fourier_grid_actual[t]:.6f}",
                    f"{wavelet_grid_actual[t]:.6f}",
                    f"{actual_solution.charge_kwh[t]:.6f}",
                    f"{actual_solution.discharge_kwh[t]:.6f}",
                    f"{actual_solution.soc_kwh[t + 1]:.6f}",
                ]
            )


def print_summary(
    labels: list[str],
    order: list[int],
    fit: FitResult,
    actual: Solution,
    fourier: Solution,
    wavelet: Solution,
    fourier_actual_cost: float,
    wavelet_actual_cost: float,
) -> None:
    grid_by_template = _to_template_order(actual.grid_kwh, order)
    requested = {
        "10:00-10:10",
        "12:00-12:10",
        "14:00-14:10",
        "16:00-16:10",
        "18:00-18:10",
        "20:00-20:10",
    }
    print("\n问题一原始数据线性规划")
    for label, value in zip(labels, grid_by_template):
        if label in requested:
            print(f"  {label}：{value:.4f} kWh")
    print(f"  全天购电量：{np.sum(actual.grid_kwh):.4f} kWh")
    print(f"  全天购电费：{actual.total_cost_yuan:.4f} 元")
    print(f"  弃光量：{np.sum(actual.curtailment_kwh):.4f} kWh")
    print(f"  SOC范围：{np.min(actual.soc_kwh):.4f}～{np.max(actual.soc_kwh):.4f} kWh")

    fourier_gap = fourier_actual_cost - actual.total_cost_yuan
    wavelet_gap = wavelet_actual_cost - actual.total_cost_yuan
    fourier_gap_rate = (
        100.0 * fourier_gap / actual.total_cost_yuan if actual.total_cost_yuan else 0.0
    )
    wavelet_gap_rate = (
        100.0 * wavelet_gap / actual.total_cost_yuan if actual.total_cost_yuan else 0.0
    )
    print("\n周期拟合、小波改进与策略对照")
    print(
        "  傅里叶负载拟合："
        f"MAE={fit.load_fourier_mae:.4f} kW，"
        f"RMSE={fit.load_fourier_rmse:.4f} kW，R^2={fit.load_fourier_r2:.6f}"
    )
    print(
        f"  傅里叶+{fit.wavelet_name}{fit.wavelet_level}层小波负载拟合："
        f"MAE={fit.load_mae:.4f} kW，RMSE={fit.load_rmse:.4f} kW，"
        f"R^2={fit.load_r2:.6f}"
    )
    print(f"  光伏拟合：MAE={fit.pv_mae:.4f} kW，RMSE={fit.pv_rmse:.4f} kW，R^2={fit.pv_r2:.6f}")
    print(
        "  光伏曲线参数："
        f"A={fit.pv_amplitude_kw:.4f} kW，日出={fit.sunrise_hour:.2f} h，"
        f"日落={fit.sunset_hour:.2f} h，γ={fit.pv_shape:.3f}"
    )
    print(f"  傅里叶场景内购电费：{fourier.total_cost_yuan:.4f} 元")
    print(f"  傅里叶策略在真实数据下费用：{fourier_actual_cost:.4f} 元")
    print(
        f"  傅里叶策略相对原始最优解成本损失："
        f"{fourier_gap:.4f} 元（{fourier_gap_rate:.6f}%）"
    )
    print(f"  小波改进场景内购电费：{wavelet.total_cost_yuan:.4f} 元")
    print(f"  小波改进策略在真实数据下费用：{wavelet_actual_cost:.4f} 元")
    print(
        f"  小波改进策略相对原始最优解成本损失："
        f"{wavelet_gap:.4f} 元（{wavelet_gap_rate:.6f}%）"
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="C题问题一线性规划与周期拟合对照")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--comparison", type=Path, default=DEFAULT_COMPARISON)
    parser.add_argument(
        "--result-scenario",
        choices=("actual", "fourier", "wavelet"),
        default="wavelet",
        help=(
            "写入结果工作簿的数据场景：actual=原始数据，"
            "fourier=傅里叶负荷，wavelet=傅里叶-小波负荷（默认）"
        ),
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    cfg = Config()
    data = read_input(args.input)
    order, labels = template_chronological_order(args.template)
    index = np.asarray(order, dtype=int)
    price = data.price[index]
    load_actual = data.load_kw[index]
    pv_actual = data.pv_kw[index]

    fit = fit_profiles(data, order, cfg)
    actual_solution = solve_linear_program(price, load_actual, pv_actual, cfg)
    fourier_solution = solve_linear_program(
        price, fit.load_fourier_kw, fit.pv_kw, cfg
    )
    wavelet_solution = solve_linear_program(price, fit.load_kw, fit.pv_kw, cfg)
    fourier_grid_actual, _, fourier_actual_cost = evaluate_fitted_dispatch(
        fourier_solution, price, load_actual, pv_actual, cfg
    )
    wavelet_grid_actual, _, wavelet_actual_cost = evaluate_fitted_dispatch(
        wavelet_solution, price, load_actual, pv_actual, cfg
    )

    scenarios = {
        "actual": actual_solution,
        "fourier": fourier_solution,
        "wavelet": wavelet_solution,
    }
    result_solution = scenarios[args.result_scenario]
    write_result(args.template, args.output, order, result_solution)
    write_comparison(
        args.comparison,
        price,
        load_actual,
        pv_actual,
        fit,
        actual_solution,
        fourier_solution,
        wavelet_solution,
        fourier_grid_actual,
        wavelet_grid_actual,
    )
    print_summary(
        labels,
        order,
        fit,
        actual_solution,
        fourier_solution,
        wavelet_solution,
        fourier_actual_cost,
        wavelet_actual_cost,
    )
    scenario_name = {
        "actual": "原始数据",
        "fourier": "傅里叶负荷+截断正弦光伏",
        "wavelet": "傅里叶-小波负荷+截断正弦光伏",
    }[args.result_scenario]
    print(f"\n正式结果（{scenario_name}）：{args.output.resolve()}")
    print(f"拟合对照：{args.comparison.resolve()}")


if __name__ == "__main__":
    main()
