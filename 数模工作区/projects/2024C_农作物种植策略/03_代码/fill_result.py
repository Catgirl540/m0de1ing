# -*- coding: utf-8 -*-
"""fill_result.py — 将 plan_case*.csv 回填竞赛结果模板 xlsx
布局：行1表头（B=地块名，C+=作物名）；第一季块（54块地）→ 第二季块（两季地块）→ 表尾注。
规则：单季作物填入第一季块；只写值，不改表格形式（模板注(3)）。
运行：python fill_result.py
"""
import csv
from pathlib import Path
import openpyxl

HERE = Path(__file__).resolve().parent
OUTDIR = HERE.parent / "06_交付"
CNAME = {}
import json
D = json.loads((HERE / "data.json").read_text(encoding="utf-8"))
CNAME = {c["id"]: c["name"] for c in D["crops"]}


def norm(s):
    return str(s).replace("\n", "").strip() if s else ""


def fill(plan_csv, template, output):
    plan = list(csv.DictReader(open(plan_csv, encoding="utf-8-sig")))
    # (year, plot, season, crop_name) -> area
    cell = {}
    for r in plan:
        season = r["季次"]
        if season == "单季":
            season = "第一季"                     # 模板注(2)
        cell[(int(r["年份"]), r["地块"].strip(), season, CNAME[int(r["作物编号"])].strip())] = float(r["面积"])

    wb = openpyxl.load_workbook(template)
    unmatched = []
    for sheet in wb.sheetnames:
        year = int(sheet)
        ws = wb[sheet]
        # 作物名 → 列号
        col_of = {}
        for c in range(3, ws.max_column + 1):
            v = norm(ws.cell(row=1, column=c).value)
            if v:
                col_of[v] = c
        season = "第一季"
        for row in range(2, ws.max_row + 1):
            label = norm(ws.cell(row=row, column=1).value)
            if label in ("第一季", "第二季"):
                season = label
            plot = ws.cell(row=row, column=2).value
            if not plot or norm(plot) == "注：":
                continue
            plot = str(plot).strip()
            for cname, col in col_of.items():
                v = cell.get((year, plot, season, cname))
                if v:
                    ws.cell(row=row, column=col).value = round(v, 2)
    # 校验：回填总面积 == 计划总面积
    wb.save(output)
    wb2 = openpyxl.load_workbook(output)
    total = 0.0
    for sheet in wb2.sheetnames:
        ws = wb2[sheet]
        for row in range(2, ws.max_row + 1):
            for col in range(3, ws.max_column + 1):
                v = ws.cell(row=row, column=col).value
                if isinstance(v, (int, float)):
                    total += v
    plan_total = sum(cell.values())
    tol = 0.005 * len(cell) + 0.05     # 901 条 × 2 位小数舍入的累计误差
    print(f"{output.name}: 回填总面积={total:,.2f} 亩 | 计划={plan_total:,.2f} 亩 | "
          f"{'一致' if abs(total - plan_total) < tol else '不一致!!'}")


def f03_rotation():
    """F03：豆类轮作示例（3 类地块各一块的 7 年季槽序列）"""
    sample = {"A1": "平旱地", "D1": "水浇地", "E1": "普通大棚"}
    for case in (1, 2):
        plan = list(csv.DictReader(open(HERE / f"plan_case{case}.csv", encoding="utf-8-sig")))
        rows = [(int(r["年份"]), r["地块"], r["季次"], CNAME[int(r["作物编号"])], float(r["面积"]))
                for r in plan if r["地块"] in sample]
        rows.sort()
        with open(HERE / f"F03_轮作示例_case{case}.csv", "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(["年份", "地块", "季次", "作物", "面积亩"])
            w.writerows(rows)
    print("F03 轮作示例已导出（A1/D1/E1）")


if __name__ == "__main__":
    OUTDIR.mkdir(exist_ok=True)
    fill(HERE / "plan_case1.csv", HERE / "结果模板/result1_1.xlsx", OUTDIR / "result1_1.xlsx")
    fill(HERE / "plan_case2.csv", HERE / "结果模板/result1_2.xlsx", OUTDIR / "result1_2.xlsx")
    f03_rotation()
