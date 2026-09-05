# -*- coding: utf-8 -*-
"""fill_paper.py — 将标题/摘要/关键词填入官方格式 Word 模板（不改动其他章节）
输入：05_论文/论文模板_官方格式.docx（new_project 复制的模板）
输出：05_论文/论文初稿_摘要填充.docx
正文各章由人按 05_论文/章节草稿.md 粘贴润色（人机分工：AI 草稿、人持笔）。
"""
from pathlib import Path
import docx

HERE = Path(__file__).resolve().parent
SRC = HERE.parent / "05_论文" / "论文模板_官方格式.docx"
DST = HERE.parent / "05_论文" / "论文初稿_摘要填充.docx"

TITLE = "基于混合整数规划与情景模拟的农作物种植策略优化研究"

ABSTRACT = [
    "本文研究华北山区某乡村 2024–2030 年农作物种植方案的优化问题。该村拥有 4 类露地"
    "1201 亩与普通、智慧大棚共 20 个，需在地块适配、禁止重茬、豆类三年轮作、种植便于"
    "管理等约束下，对 41 种作物逐年逐季安排种植面积，使七年总净收益最大。",
    "针对问题一，建立多季种植计划的混合整数规划（MILP）模型，将超产部分“滞销”与“五折"
    "甩卖”两种情形统一为带销量上限的线性收益结构，采用 HiGHS 求解器精确求解，得七年总"
    "净收益分别为 7 816.4 万元与 8 651.4 万元，五折情形比滞销情形多回收 835.0 万元"
    "（+10.7%），方案已逐年逐季回填 result1_1 与 result1_2 表。",
    "针对问题二，按题给区间构造销量、成本、价格的中枢变化路径并重新求解，得净收益"
    " 9 402.1 万元（方案回填 result2）；20 组均匀抽样情景检验表明，沿用该方案的收益均值为"
    " 9 134.5 万元、变异系数仅 1.7%，按情景重优化平均仅再提高 0.2%，方案对参数波动稳健。",
    "针对问题三，引入蔬菜类需求的共同市场因子与需求—价格负相关机制，进行 8 组相关蒙特"
    "卡洛模拟并逐组重优化：重优化平均净收益 14 041.1 万元，比沿用问题二方案提高 1 327.5 "
    "万元（+10.4%），表明存在相关冲击时应逐年滚动优化。",
    "本文的亮点在于：约束体系完整覆盖题目全部硬性要求且线性可核验；“中枢求解—情景评估—"
    "重优化对比”三层递进量化了不确定性与灵活调整的价值；全部结果可由附录命令一键复现。",
]
KEYWORDS = "农作物种植策略；混合整数规划；销量线性化；情景模拟；轮作约束"


def set_text(par, text):
    """保留段落样式，仅替换文本（改第一个 run，删其余）"""
    if not par.runs:
        par.add_run(text)
        return
    par.runs[0].text = text
    for r in par.runs[1:]:
        r.text = ""


def main():
    doc = docx.Document(str(SRC))
    done = set()
    for par in doc.paragraphs:
        t = par.text.strip()
        if t.startswith("【论文标题"):
            set_text(par, TITLE); done.add("title")
        elif t.startswith("【提示】五要素"):
            set_text(par, ""); done.add("hint1")
        elif t.startswith("本文研究"):
            set_text(par, ABSTRACT[0]); done.add("a1")
        elif t.startswith("针对问题一"):
            set_text(par, ABSTRACT[1]); done.add("a2")
        elif t.startswith("针对问题二"):
            set_text(par, ABSTRACT[2]); done.add("a3")
        elif t.startswith("针对问题三"):
            set_text(par, ABSTRACT[3]); done.add("a4")
        elif t.startswith("本文的亮点在于"):
            set_text(par, ABSTRACT[4]); done.add("a5")
        elif "【关键词1" in t:
            for r in par.runs:
                if "【关键词1" in r.text:
                    r.text = KEYWORDS
            done.add("kw")
    print("已填充：", sorted(done))
    missing = {"title", "a1", "a2", "a3", "a4", "a5", "kw"} - done
    if missing:
        raise SystemExit(f"未匹配到占位段：{missing}（模板结构可能已变动）")
    doc.save(str(DST))
    print("OK ->", DST)


if __name__ == "__main__":
    main()
