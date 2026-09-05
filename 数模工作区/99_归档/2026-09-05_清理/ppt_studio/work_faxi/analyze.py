# -*- coding: utf-8 -*-
"""Dump structure/style of the source pptx for analysis."""
import sys, io, json
from pptx import Presentation
from pptx.util import Emu

SRC = r"C:\Users\yan\Desktop\发析生物ppt0901.pptx"

def emu2in(v):
    return round(Emu(v).inches, 2) if v is not None else None

def color_of(obj):
    try:
        c = obj.fill
        if c.type is not None and str(c.type) != 'MSO_FILL_TYPE.BACKGROUND (5)':
            if c.type == 1:  # solid
                col = c.fore_color
                try:
                    return str(col.rgb)
                except Exception:
                    return 'theme:' + str(col.theme_color)
            return str(c.type)
    except Exception:
        pass
    return None

def font_of(run):
    f = run.font
    col = None
    try:
        if f.color and f.color.type is not None:
            try:
                col = str(f.color.rgb)
            except Exception:
                col = 'theme:' + str(f.color.type)
    except Exception:
        col = None
    return {"text": run.text, "size": f.size.pt if f.size else None, "bold": f.bold,
            "name": f.name, "color": col}

def walk(shapes, indent=0):
    out = []
    for s in shapes:
        t = s.shape_type
        item = {"type": str(t), "name": s.name, "id": s.shape_id,
                "pos": [emu2in(s.left), emu2in(s.top), emu2in(s.width), emu2in(s.height)]}
        if s.shape_type == 6:  # group
            item["children"] = walk(s.shapes, indent+1)
        elif s.has_text_frame:
            paras = []
            for p in s.text_frame.paragraphs:
                runs = [font_of(r) for r in p.runs]
                if runs or p.text:
                    paras.append({"align": str(p.alignment), "runs": runs})
            if paras:
                item["paras"] = paras
        try:
            fc = color_of(s)
            if fc:
                item["fill"] = fc
        except Exception:
            pass
        if s.has_text_frame is False and 'GRAPHIC' not in str(t):
            pass
        out.append(item)
    return out

prs = Presentation(SRC)
print(f"SLIDE SIZE: {emu2in(prs.slide_width)} x {emu2in(prs.slide_height)} in")
print(f"SLIDES: {len(prs.slides)}")
data = []
for i, slide in enumerate(prs.slides):
    sl = {"idx": i, "layout": slide.slide_layout.name, "shapes": walk(slide.shapes)}
    data.append(sl)
    print(f"\n=== SLIDE {i} (layout={sl['layout']}) ===")
    print(json.dumps(sl["shapes"], ensure_ascii=False, indent=None))
