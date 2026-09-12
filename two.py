"""2026 CUMCM C题问题二：滚动傅里叶—小波—正弦预测与购电线性规划。

每天 0:00 仅使用此前已观测的数据预测当天负荷和光伏功率：
1. 年度谐波—星期效应岭回归与历史轮廓融合预测负荷；
2. 二阶傅里叶与 Haar 小波重构日内负荷残差；
3. 截断幂正弦函数拟合光伏功率；
4. 历史净负荷预测误差的 80% 分位数作为安全裕度；
5. 采用 48 小时滚动线性规划，执行前 24 小时并连续传递 SOC；
6. 用真实数据回放，供电不足部分按当时电价 5 倍紧急购电。
7. 采用“训训训验训测”循环划分，并用严格前向贝叶斯优化和早停选择负荷参数。

默认输入：C题/附件/csv/附件1.csv 与附件2的两个 CSV
默认模板：C题/附件/附件5/result2.xlsx
默认输出：C题/result2.xlsx
默认汇总：C题/问题二逐日汇总.csv
默认检验：C题/问题二_训练验证测试检验.csv

依赖：numpy、scipy、openpyxl
正式求解：python two.py
贝叶斯寻优：python two.py --bayes-optimize
自定义寻优：python two.py --bayes-optimize --bayes-trials 50 --bayes-patience 12 --bayes-seed 2026
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import os
import shutil
from copy import copy
from dataclasses import asdict, dataclass, replace
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
DEFAULT_BASELINE = ROOT / "C题" / "附件" / "csv" / "附件1.csv"
DEFAULT_LOAD = ROOT / "C题" / "附件" / "csv" / "附件2_小区负载.csv"
DEFAULT_PV = ROOT / "C题" / "附件" / "csv" / "附件2_光伏发电实际功率.csv"
DEFAULT_TEMPLATE = ROOT / "C题" / "附件" / "附件5" / "result2.xlsx"
DEFAULT_OUTPUT = ROOT / "C题" / "result2.xlsx"
DEFAULT_SUMMARY = ROOT / "C题" / "问题二逐日汇总.csv"
DEFAULT_SPLIT_REPORT = ROOT / "C题" / "问题二_训练验证测试检验.csv"
DEFAULT_BAYES_REPORT = ROOT / "C题" / "贝叶斯优化_负荷参数与评价_严格前向.json"


@dataclass(frozen=True)
class Config:
    dt_hours: float = 1.0 / 6.0
    soc_min_kwh: float = 1200.0
    soc_max_kwh: float = 10800.0
    initial_soc_kwh: float = 6000.0
    max_charge_power_kw: float = 5000.0
    max_discharge_power_kw: float = 5000.0
    charge_efficiency: float = 0.90
    discharge_efficiency: float = 0.90
    history_days: int = 28
    risk_days: int = 28
    minimum_risk_days: int = 7
    half_life_days: float = 7.0
    same_weekday_multiplier: float = 2.0
    wavelet_level: int = 3
    annual_harmonics: int = 3
    load_ridge_lambda: float = 3.0
    dynamic_load_blend: float = 0.65
    use_bayesian_load_schedule: bool = True
    risk_quantile: float = 0.80
    emergency_price_multiplier: float = 5.0
    throughput_tie_breaker: float = 1.0e-7
    emergency_tolerance_kwh: float = 1.0e-5

    @property
    def max_charge_kwh(self) -> float:
        return self.max_charge_power_kw * self.dt_hours

    @property
    def max_discharge_kwh(self) -> float:
        return self.max_discharge_power_kw * self.dt_hours


# 严格前向贝叶斯优化得到的参数轨迹。每组参数只从生效日前已经完成的
# 验证日获得，因而不会让未来验证值或测试值进入当前预测。日期表示参数
# 最早可用于0:00计划的日期；未到首个日期时沿用 Config 的原始参数。
BAYESIAN_LOAD_PARAMETER_SCHEDULE: tuple[
    tuple[date, dict[str, float | int]], ...
] = (
    (
        date(2025, 2, 16),
        {
            "history_days": 38,
            "half_life_days": 25.29917197614039,
            "same_weekday_multiplier": 2.241902727000516,
            "wavelet_level": 2,
            "annual_harmonics": 3,
            "load_ridge_lambda": 0.003840001534395503,
            "dynamic_load_blend": 0.6206534746761987,
        },
    ),
    (
        date(2025, 3, 12),
        {
            "history_days": 27,
            "half_life_days": 27.866925493061643,
            "same_weekday_multiplier": 2.0302869528791616,
            "wavelet_level": 2,
            "annual_harmonics": 3,
            "load_ridge_lambda": 0.002040174751235714,
            "dynamic_load_blend": 0.6470383902370724,
        },
    ),
    (
        date(2025, 4, 23),
        {
            "history_days": 35,
            "half_life_days": 22.86236550087216,
            "same_weekday_multiplier": 1.7054409173409268,
            "wavelet_level": 2,
            "annual_harmonics": 3,
            "load_ridge_lambda": 0.0010398473604926101,
            "dynamic_load_blend": 0.8157826442179166,
        },
    ),
    (
        date(2025, 5, 17),
        {
            "history_days": 23,
            "half_life_days": 22.16005337031691,
            "same_weekday_multiplier": 3.4204172241648596,
            "wavelet_level": 2,
            "annual_harmonics": 3,
            "load_ridge_lambda": 0.00010683581757667908,
            "dynamic_load_blend": 0.726727237291571,
        },
    ),
    (
        date(2025, 8, 3),
        {
            "history_days": 35,
            "half_life_days": 22.86236550087216,
            "same_weekday_multiplier": 1.7054409173409268,
            "wavelet_level": 2,
            "annual_harmonics": 3,
            "load_ridge_lambda": 0.0010398473604926101,
            "dynamic_load_blend": 0.8157826442179166,
        },
    ),
)


def bayesian_load_config(cfg: Config, issue_date: date) -> Config:
    """返回计划发布日可合法使用的已冻结负荷参数。"""
    if not cfg.use_bayesian_load_schedule:
        return cfg
    overrides: dict[str, float | int] | None = None
    for effective_date, candidate in BAYESIAN_LOAD_PARAMETER_SCHEDULE:
        if issue_date < effective_date:
            break
        overrides = candidate
    return cfg if overrides is None else replace(cfg, **overrides)


@dataclass(frozen=True)
class DatasetSplit:
    """时间序列数据集边界；验证集和测试集绝不参与提前拟合。"""

    train_end: date = date(2025, 7, 31)
    validation_end: date = date(2025, 9, 30)

    def training_cutoff(self, dates: list[date]) -> int:
        """返回训练集结束后的首个索引，供后续阶段冻结拟合样本。"""
        if self.validation_end <= self.train_end:
            raise ValueError("验证集结束日期必须晚于训练集结束日期。")
        cutoff = sum(day <= self.train_end for day in dates)
        if cutoff <= 0 or cutoff >= len(dates):
            raise ValueError("训练集必须非空，且不能覆盖全部数据。")
        return cutoff

    def label(self, day: date) -> str:
        if day <= self.train_end:
            return "训练集"
        if day <= self.validation_end:
            return "验证集"
        return "测试集"

    def cycle_labels(self, dates: list[date]) -> dict[str, list[int]]:
        """全年按固定循环均匀抽取：训练、训练、训练、验证、训练、测试。"""
        groups = {"训练集": [], "验证集": [], "测试集": []}
        for index, _ in enumerate(dates):
            slot = index % 6
            label = "验证集" if slot == 3 else "测试集" if slot == 5 else "训练集"
            groups[label].append(index)
        return groups


@dataclass
class BaselineData:
    price: np.ndarray
    load_kw: np.ndarray
    pv_kw: np.ndarray


@dataclass
class HistoricalData:
    dates: list[date]
    load_kw: np.ndarray
    pv_kw: np.ndarray


@dataclass
class Forecast:
    load_kw: np.ndarray
    pv_kw: np.ndarray
    load_fourier_kw: np.ndarray
    pv_amplitude_kw: float
    sunrise_hour: float
    sunset_hour: float
    pv_shape: float


@dataclass
class LPSolution:
    grid_kwh: np.ndarray
    charge_kwh: np.ndarray
    discharge_kwh: np.ndarray
    curtailment_kwh: np.ndarray
    soc_kwh: np.ndarray
    objective_yuan: float


@dataclass
class DailyResult:
    day: date
    forecast_load_kw: np.ndarray
    forecast_pv_kw: np.ndarray
    reserve_kwh: np.ndarray
    grid_kwh: np.ndarray
    charge_kwh: np.ndarray
    discharge_kwh: np.ndarray
    planned_curtailment_kwh: np.ndarray
    actual_curtailment_kwh: np.ndarray
    emergency_kwh: np.ndarray
    soc_start_kwh: float
    soc_end_kwh: float
    plan_cost_yuan: float
    emergency_cost_yuan: float

    @property
    def total_cost_yuan(self) -> float:
        return self.plan_cost_yuan + self.emergency_cost_yuan


def _float(value: object, context: str) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{context}不是数值：{value!r}") from exc
    if not math.isfinite(number):
        raise ValueError(f"{context}不是有限数值：{value!r}")
    return number


def _parse_source_minute(value: str) -> int:
    text = value.strip().replace("+1", "")
    parts = text.split(":")
    if len(parts) < 2:
        raise ValueError(f"无法解析时间：{value!r}")
    hour, minute = int(parts[0]) % 24, int(parts[1])
    return hour * 60 + minute


def read_baseline(path: Path) -> tuple[BaselineData, list[int]]:
    if not path.exists():
        raise FileNotFoundError(f"找不到附件1 CSV：{path}")
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        rows = list(csv.reader(file))
    if len(rows) != 145 or len(rows[0]) < 4:
        raise ValueError("附件1应包含1行表头和144行数据。")

    times: list[str] = []
    price: list[float] = []
    load: list[float] = []
    pv: list[float] = []
    for row_number, row in enumerate(rows[1:], start=2):
        times.append(row[0].strip())
        price.append(_float(row[1], f"附件1第{row_number}行电价"))
        load.append(_float(row[2], f"附件1第{row_number}行负荷"))
        pv.append(_float(row[3], f"附件1第{row_number}行光伏"))

    minutes = [_parse_source_minute(value) for value in times]
    if sorted(minutes) != list(range(0, 1440, 10)):
        raise ValueError("附件1没有完整覆盖一天的144个10分钟时刻。")
    chronological_order = sorted(range(144), key=minutes.__getitem__)
    order = np.asarray(chronological_order, dtype=int)
    return (
        BaselineData(
            np.asarray(price, dtype=float)[order],
            np.asarray(load, dtype=float)[order],
            np.asarray(pv, dtype=float)[order],
        ),
        chronological_order,
    )


def read_daily_matrix(path: Path, chronological_order: list[int]) -> tuple[list[date], np.ndarray]:
    if not path.exists():
        raise FileNotFoundError(f"找不到附件2 CSV：{path}")
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.reader(file)
        header = next(reader, None)
        rows = list(reader)
    if header is None or len(header) != 145:
        raise ValueError(f"{path.name}应包含日期列和144个时刻列。")
    if len(rows) != 365:
        raise ValueError(f"{path.name}应包含2025年的365天数据。")

    dates: list[date] = []
    values: list[list[float]] = []
    for row_number, row in enumerate(rows, start=2):
        if len(row) != 145:
            raise ValueError(f"{path.name}第{row_number}行列数不是145。")
        try:
            current_date = datetime.fromisoformat(row[0].strip()).date()
        except ValueError as exc:
            raise ValueError(f"{path.name}第{row_number}行日期无法解析：{row[0]!r}") from exc
        dates.append(current_date)
        values.append(
            [_float(row[i + 1], f"{path.name}第{row_number}行第{i + 2}列") for i in chronological_order]
        )

    matrix = np.asarray(values, dtype=float)
    if np.any(matrix < 0):
        raise ValueError(f"{path.name}中存在负功率。")
    return dates, matrix


def read_historical_data(
    load_path: Path, pv_path: Path, chronological_order: list[int]
) -> HistoricalData:
    load_dates, load = read_daily_matrix(load_path, chronological_order)
    pv_dates, pv = read_daily_matrix(pv_path, chronological_order)
    if load_dates != pv_dates:
        raise ValueError("附件2的负荷日期和光伏日期不一致。")
    expected = [date(2025, 1, 1) + timedelta(days=i) for i in range(365)]
    if load_dates != expected:
        raise ValueError("附件2日期应连续覆盖2025-01-01至2025-12-31。")
    return HistoricalData(load_dates, load, pv)


def _fourier_design(hours: np.ndarray) -> np.ndarray:
    omega = 2.0 * np.pi / 24.0
    return np.column_stack(
        [
            np.ones_like(hours),
            np.cos(omega * hours),
            np.sin(omega * hours),
            np.cos(2.0 * omega * hours),
            np.sin(2.0 * omega * hours),
        ]
    )


def fit_load_fourier(hours: np.ndarray, load_kw: np.ndarray) -> np.ndarray:
    design = _fourier_design(hours)
    coefficients, *_ = np.linalg.lstsq(design, load_kw, rcond=None)
    return np.maximum(0.0, design @ coefficients)


def _haar_dwt(signal: np.ndarray, level: int) -> tuple[np.ndarray, list[np.ndarray]]:
    approximation = np.asarray(signal, dtype=float).copy()
    details: list[np.ndarray] = []
    scale = math.sqrt(2.0)
    for _ in range(level):
        if approximation.size % 2:
            raise ValueError("Haar 小波每层输入长度必须为偶数。")
        even = approximation[0::2]
        odd = approximation[1::2]
        details.append((even - odd) / scale)
        approximation = (even + odd) / scale
    return approximation, details


def _haar_idwt(approximation: np.ndarray, details: list[np.ndarray]) -> np.ndarray:
    reconstructed = np.asarray(approximation, dtype=float).copy()
    scale = math.sqrt(2.0)
    for detail in reversed(details):
        restored = np.empty(reconstructed.size * 2, dtype=float)
        restored[0::2] = (reconstructed + detail) / scale
        restored[1::2] = (reconstructed - detail) / scale
        reconstructed = restored
    return reconstructed


def wavelet_correct_residual(
    profile_kw: np.ndarray, fourier_kw: np.ndarray, level: int
) -> np.ndarray:
    residual = np.asarray(profile_kw) - np.asarray(fourier_kw)
    if residual.size % (2**level):
        raise ValueError(f"样本数必须能被2^{level}整除。")
    approximation, details = _haar_dwt(residual, level)
    thresholded: list[np.ndarray] = []
    for detail in details:
        sigma = float(np.median(np.abs(detail - np.median(detail))) / 0.6745)
        threshold = sigma * math.sqrt(2.0 * math.log(residual.size))
        thresholded.append(
            np.sign(detail) * np.maximum(np.abs(detail) - threshold, 0.0)
        )
    return _haar_idwt(approximation, thresholded)


def fit_pv_truncated_sine(
    hours: np.ndarray, pv_kw: np.ndarray
) -> tuple[np.ndarray, float, float, float, float]:
    """用小规模网格搜索拟合截断幂正弦函数。"""
    peak = float(np.max(pv_kw))
    if peak <= 1.0e-8:
        return np.zeros_like(pv_kw), 0.0, 0.0, 0.0, 1.0

    positive = np.flatnonzero(pv_kw > max(1.0, peak * 0.002))
    if positive.size == 0:
        return np.zeros_like(pv_kw), 0.0, 0.0, 0.0, 1.0
    first_hour = float(hours[positive[0]])
    last_hour = float(hours[positive[-1]])
    sunrise_grid = np.arange(max(0.0, first_hour - 0.5), first_hour + 1.0e-9, 1.0 / 6.0)
    sunset_grid = np.arange(last_hour, min(24.0, last_hour + 0.5) + 1.0e-9, 1.0 / 6.0)
    shape_grid = np.linspace(0.6, 2.8, 45)

    best_sse = math.inf
    best_curve = np.zeros_like(pv_kw)
    best_parameters = (0.0, first_hour, last_hour, 1.0)
    for sunrise in sunrise_grid:
        for sunset in sunset_grid:
            if sunset - sunrise < 5.0:
                continue
            phase = np.pi * (hours - sunrise) / (sunset - sunrise)
            sine = np.zeros_like(hours)
            daylight = (hours >= sunrise) & (hours <= sunset)
            sine[daylight] = np.maximum(0.0, np.sin(phase[daylight]))
            basis = sine[None, :] ** shape_grid[:, None]
            denominator = np.sum(basis * basis, axis=1)
            numerator = basis @ pv_kw
            amplitudes = np.divide(
                numerator,
                denominator,
                out=np.zeros_like(numerator),
                where=denominator > 0,
            )
            amplitudes = np.maximum(0.0, amplitudes)
            curves = amplitudes[:, None] * basis
            errors = np.sum((curves - pv_kw[None, :]) ** 2, axis=1)
            index = int(np.argmin(errors))
            if float(errors[index]) < best_sse:
                best_sse = float(errors[index])
                best_curve = curves[index].copy()
                best_parameters = (
                    float(amplitudes[index]),
                    float(sunrise),
                    float(sunset),
                    float(shape_grid[index]),
                )
    return best_curve, *best_parameters


def _weighted_history_profile(
    matrix: np.ndarray,
    historical_dates: list[date],
    cutoff: int,
    target_date: date,
    baseline: np.ndarray,
    cfg: Config,
    allowed_indices: set[int] | None = None,
) -> np.ndarray:
    if cutoff <= 0:
        return baseline.copy()
    first = max(0, cutoff - cfg.history_days)
    indices = np.arange(first, cutoff, dtype=int)
    if allowed_indices is not None:
        indices = np.asarray([i for i in indices if i in allowed_indices], dtype=int)
    if indices.size == 0:
        return baseline.copy()
    ages = cutoff - 1 - indices
    weights = np.exp(-math.log(2.0) * ages / cfg.half_life_days)
    same_weekday = np.asarray(
        [historical_dates[i].weekday() == target_date.weekday() for i in indices],
        dtype=float,
    )
    weights *= 1.0 + cfg.same_weekday_multiplier * same_weekday
    profile = np.average(matrix[indices], axis=0, weights=weights)

    # 前7天历史样本较少，逐步降低附件1先验曲线的权重。
    prior_share = max(0.0, (7.0 - cutoff) / 7.0)
    return (1.0 - prior_share) * profile + prior_share * baseline


def _calendar_features(day: date, harmonics: int) -> np.ndarray:
    """年度谐波与星期效应；全部特征在预测日前即可确定。"""
    day_of_year = day.timetuple().tm_yday - 1
    columns = [1.0]
    for harmonic in range(1, harmonics + 1):
        angle = 2.0 * math.pi * harmonic * day_of_year / 365.0
        columns.extend([math.cos(angle), math.sin(angle)])
    weekday = day.weekday()
    columns.extend([1.0 if weekday == value else 0.0 for value in range(1, 7)])
    return np.asarray(columns, dtype=float)


def fit_dynamic_load_profile(
    historical: HistoricalData,
    cutoff: int,
    target_date: date,
    baseline: np.ndarray,
    cfg: Config,
    allowed_indices: set[int] | None,
) -> np.ndarray:
    """用训练日的年度周期和星期效应岭回归预测完整日负荷曲线。"""
    indices = [
        index for index in range(cutoff)
        if allowed_indices is None or index in allowed_indices
    ]
    feature_count = 1 + 2 * cfg.annual_harmonics + 6
    if len(indices) < max(14, feature_count + 1):
        return baseline.copy()
    design = np.vstack(
        [_calendar_features(historical.dates[index], cfg.annual_harmonics) for index in indices]
    )
    target = historical.load_kw[indices]
    penalty = np.eye(feature_count, dtype=float) * cfg.load_ridge_lambda
    penalty[0, 0] = 0.0
    coefficients = np.linalg.solve(
        design.T @ design + penalty,
        design.T @ target,
    )
    prediction = _calendar_features(target_date, cfg.annual_harmonics) @ coefficients
    return np.maximum(0.0, prediction)


def make_forecast(
    historical: HistoricalData,
    baseline: BaselineData,
    cutoff: int,
    target_date: date,
    hours: np.ndarray,
    cfg: Config,
    allowed_indices: set[int] | None = None,
    pv_cfg: Config | None = None,
) -> Forecast:
    pv_settings = cfg if pv_cfg is None else pv_cfg
    load_profile = _weighted_history_profile(
        historical.load_kw,
        historical.dates,
        cutoff,
        target_date,
        baseline.load_kw,
        cfg,
        allowed_indices,
    )
    dynamic_load = fit_dynamic_load_profile(
        historical,
        cutoff,
        target_date,
        baseline.load_kw,
        cfg,
        allowed_indices,
    )
    blend = float(np.clip(cfg.dynamic_load_blend, 0.0, 1.0))
    load_profile = (1.0 - blend) * load_profile + blend * dynamic_load
    pv_profile = _weighted_history_profile(
        historical.pv_kw,
        historical.dates,
        cutoff,
        target_date,
        baseline.pv_kw,
        pv_settings,
        allowed_indices,
    )

    load_fourier = fit_load_fourier(hours, load_profile)
    wavelet_residual = wavelet_correct_residual(
        load_profile, load_fourier, cfg.wavelet_level
    )
    load_forecast = np.maximum(0.0, load_fourier + wavelet_residual)
    pv_forecast, amplitude, sunrise, sunset, shape = fit_pv_truncated_sine(
        hours, pv_profile
    )
    return Forecast(
        load_forecast,
        pv_forecast,
        load_fourier,
        amplitude,
        sunrise,
        sunset,
        shape,
    )


def risk_reserve(error_history: list[np.ndarray], cfg: Config) -> np.ndarray:
    if len(error_history) < cfg.minimum_risk_days:
        return np.zeros(144, dtype=float)
    errors = np.asarray(error_history[-cfg.risk_days :], dtype=float)
    reserve = np.quantile(errors, cfg.risk_quantile, axis=0)
    return np.maximum(0.0, reserve)


def solve_rolling_lp(
    price: np.ndarray,
    load_kw: np.ndarray,
    pv_kw: np.ndarray,
    initial_soc_kwh: float,
    cfg: Config,
) -> LPSolution:
    try:
        from scipy.optimize import linprog
        from scipy.sparse import coo_matrix
    except ModuleNotFoundError as exc:
        raise ModuleNotFoundError("缺少scipy，请先执行：python -m pip install scipy") from exc

    n = len(price)
    if n not in (144, 288) or len(load_kw) != n or len(pv_kw) != n:
        raise ValueError("滚动优化数组长度必须为144或288且彼此一致。")
    load_energy = load_kw * cfg.dt_hours
    pv_energy = pv_kw * cfg.dt_hours

    g0, c0, d0, w0, e0 = 0, n, 2 * n, 3 * n, 4 * n
    variable_count = 5 * n + 1
    objective = np.zeros(variable_count, dtype=float)
    objective[g0 : g0 + n] = price
    objective[c0 : c0 + n] = cfg.throughput_tie_breaker
    objective[d0 : d0 + n] = cfg.throughput_tie_breaker

    rows: list[int] = []
    columns: list[int] = []
    coefficients: list[float] = []
    b_eq = np.zeros(2 * n, dtype=float)
    for t in range(n):
        # g + PV + d = load + c + w
        rows.extend([t, t, t, t])
        columns.extend([g0 + t, c0 + t, d0 + t, w0 + t])
        coefficients.extend([1.0, -1.0, 1.0, -1.0])
        b_eq[t] = load_energy[t] - pv_energy[t]

        # e[t+1] = e[t] + eta_c*c - d/eta_d
        row = n + t
        rows.extend([row, row, row, row])
        columns.extend([c0 + t, d0 + t, e0 + t, e0 + t + 1])
        coefficients.extend(
            [-cfg.charge_efficiency, 1.0 / cfg.discharge_efficiency, -1.0, 1.0]
        )

    a_eq = coo_matrix(
        (coefficients, (rows, columns)), shape=(2 * n, variable_count)
    ).tocsr()

    bounds: list[tuple[float | None, float | None]] = []
    bounds.extend([(0.0, None)] * n)
    bounds.extend([(0.0, cfg.max_charge_kwh)] * n)
    bounds.extend([(0.0, cfg.max_discharge_kwh)] * n)
    bounds.extend([(0.0, None)] * n)
    bounds.extend([(cfg.soc_min_kwh, cfg.soc_max_kwh)] * (n + 1))
    bounds[e0] = (initial_soc_kwh, initial_soc_kwh)
    # 48小时末恢复当前SOC，避免滚动窗口末端无成本放空储能。
    bounds[e0 + n] = (initial_soc_kwh, initial_soc_kwh)

    result = linprog(
        objective,
        A_eq=a_eq,
        b_eq=b_eq,
        bounds=bounds,
        method="highs",
        options={"presolve": True},
    )
    if not result.success:
        raise RuntimeError(f"线性规划失败：status={result.status}, {result.message}")

    x = np.asarray(result.x, dtype=float)
    solution = LPSolution(
        x[g0 : g0 + n],
        x[c0 : c0 + n],
        x[d0 : d0 + n],
        x[w0 : w0 + n],
        x[e0 : e0 + n + 1],
        float(np.dot(price, x[g0 : g0 + n])),
    )
    validate_lp(solution, load_energy, pv_energy, initial_soc_kwh, cfg)
    return solution


def validate_lp(
    solution: LPSolution,
    load_energy: np.ndarray,
    pv_energy: np.ndarray,
    initial_soc_kwh: float,
    cfg: Config,
    tolerance: float = 1.0e-5,
) -> None:
    balance = (
        solution.grid_kwh
        + pv_energy
        + solution.discharge_kwh
        - load_energy
        - solution.charge_kwh
        - solution.curtailment_kwh
    )
    next_soc = (
        solution.soc_kwh[:-1]
        + cfg.charge_efficiency * solution.charge_kwh
        - solution.discharge_kwh / cfg.discharge_efficiency
    )
    checks = {
        "电量平衡": float(np.max(np.abs(balance))),
        "SOC递推": float(np.max(np.abs(solution.soc_kwh[1:] - next_soc))),
        "初始SOC": abs(float(solution.soc_kwh[0]) - initial_soc_kwh),
        "滚动窗口终止SOC": abs(float(solution.soc_kwh[-1]) - initial_soc_kwh),
    }
    failed = {name: value for name, value in checks.items() if value > tolerance}
    if failed:
        raise RuntimeError(f"线性规划结果校验失败：{failed}")
    if np.min(solution.soc_kwh) < cfg.soc_min_kwh - tolerance:
        raise RuntimeError("SOC低于安全下限。")
    if np.max(solution.soc_kwh) > cfg.soc_max_kwh + tolerance:
        raise RuntimeError("SOC高于安全上限。")
    if np.max(np.minimum(solution.charge_kwh, solution.discharge_kwh)) > 1.0e-4:
        raise RuntimeError("线性规划结果出现同一时段同时充放电。")


def replay_actual_day(
    actual_load_kw: np.ndarray,
    actual_pv_kw: np.ndarray,
    grid_kwh: np.ndarray,
    charge_kwh: np.ndarray,
    discharge_kwh: np.ndarray,
    cfg: Config,
) -> tuple[np.ndarray, np.ndarray]:
    deficit = (
        actual_load_kw * cfg.dt_hours
        + charge_kwh
        - actual_pv_kw * cfg.dt_hours
        - discharge_kwh
        - grid_kwh
    )
    emergency = np.maximum(0.0, deficit)
    curtailment = np.maximum(0.0, -deficit)
    balance = (
        grid_kwh
        + actual_pv_kw * cfg.dt_hours
        + discharge_kwh
        + emergency
        - actual_load_kw * cfg.dt_hours
        - charge_kwh
        - curtailment
    )
    if float(np.max(np.abs(balance))) > 1.0e-5:
        raise RuntimeError("真实数据回放电量不平衡。")
    return emergency, curtailment


def simulate_year(
    baseline: BaselineData,
    historical: HistoricalData,
    cfg: Config,
    training_cutoff: int | None = None,
    training_indices: set[int] | None = None,
) -> list[DailyResult]:
    hours = np.arange(144, dtype=float) * cfg.dt_hours
    price_48 = np.tile(baseline.price, 2)
    error_history: list[np.ndarray] = []
    results: list[DailyResult] = []
    soc = cfg.initial_soc_kwh

    for day_index, current_date in enumerate(historical.dates):
        # 训练期采用因果滚动拟合；进入验证/测试期后冻结训练样本边界。
        model_cutoff = (
            day_index
            if training_cutoff is None
            else min(day_index, training_cutoff)
        )
        forecast_cfg = bayesian_load_config(cfg, current_date)
        current_forecast = make_forecast(
            historical, baseline, model_cutoff, current_date, hours, forecast_cfg,
            allowed_indices=training_indices,
            pv_cfg=cfg,
        )
        next_forecast = make_forecast(
            historical,
            baseline,
            model_cutoff,
            current_date + timedelta(days=1),
            hours,
            forecast_cfg,
            allowed_indices=training_indices,
            pv_cfg=cfg,
        )
        reserve = risk_reserve(error_history, cfg)
        risk_load_current = current_forecast.load_kw + reserve / cfg.dt_hours
        risk_load_next = next_forecast.load_kw + reserve / cfg.dt_hours

        solution = solve_rolling_lp(
            price_48,
            np.concatenate([risk_load_current, risk_load_next]),
            np.concatenate([current_forecast.pv_kw, next_forecast.pv_kw]),
            soc,
            cfg,
        )
        grid = solution.grid_kwh[:144]
        charge = solution.charge_kwh[:144]
        discharge = solution.discharge_kwh[:144]
        planned_curtailment = solution.curtailment_kwh[:144]
        emergency, actual_curtailment = replay_actual_day(
            historical.load_kw[day_index],
            historical.pv_kw[day_index],
            grid,
            charge,
            discharge,
            cfg,
        )

        soc_start = soc
        soc = float(solution.soc_kwh[144])
        plan_cost = float(np.dot(baseline.price, grid))
        emergency_cost = float(
            np.dot(
                cfg.emergency_price_multiplier * baseline.price,
                emergency,
            )
        )
        results.append(
            DailyResult(
                current_date,
                current_forecast.load_kw,
                current_forecast.pv_kw,
                reserve,
                grid,
                charge,
                discharge,
                planned_curtailment,
                actual_curtailment,
                emergency,
                soc_start,
                soc,
                plan_cost,
                emergency_cost,
            )
        )

        actual_net = (
            historical.load_kw[day_index] - historical.pv_kw[day_index]
        ) * cfg.dt_hours
        forecast_net = (
            current_forecast.load_kw - current_forecast.pv_kw
        ) * cfg.dt_hours
        if training_indices is None or day_index in training_indices:
            error_history.append(actual_net - forecast_net)

        if (day_index + 1) % 30 == 0 or day_index == len(historical.dates) - 1:
            print(
                f"已完成 {current_date.isoformat()}："
                f"SOC={soc:.2f} kWh，紧急购电={np.sum(emergency):.2f} kWh"
            )
    return results


def _to_template_order(values: np.ndarray, chronological_order: list[int]) -> np.ndarray:
    mapped = np.empty_like(values)
    mapped[np.asarray(chronological_order, dtype=int)] = values
    return mapped


def _copy_cell_style(source, target) -> None:
    target._style = copy(source._style)
    target.number_format = source.number_format
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)


def _interval_label(start_index: int, end_index_exclusive: int) -> str:
    start_minutes = start_index * 10
    end_minutes = end_index_exclusive * 10

    def format_minutes(minutes: int) -> str:
        if minutes == 1440:
            return "24:00"
        return f"{minutes // 60}:{minutes % 60:02d}"

    return f"{format_minutes(start_minutes)}-{format_minutes(end_minutes)}"


def group_emergency(values: np.ndarray, tolerance: float) -> list[tuple[str, float]]:
    positive = values > tolerance
    groups: list[tuple[str, float]] = []
    index = 0
    while index < len(values):
        if not positive[index]:
            index += 1
            continue
        start = index
        while index + 1 < len(values) and positive[index + 1]:
            index += 1
        stop = index + 1
        groups.append((_interval_label(start, stop), float(np.sum(values[start:stop]))))
        index += 1
    return groups


def write_result_workbook(
    template: Path,
    output: Path,
    results: list[DailyResult],
    price: np.ndarray,
    chronological_order: list[int],
    cfg: Config,
) -> None:
    if not template.exists():
        raise FileNotFoundError(f"找不到result2模板：{template}")
    if template.resolve() == output.resolve():
        raise ValueError("输出文件不能覆盖官方模板。")
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(f".{output.stem}.building{output.suffix}")
    if temporary.exists():
        temporary.unlink()
    shutil.copy2(template, temporary)

    try:
        workbook = load_workbook(temporary)
        plan_sheet = workbook["计划购电量"]
        storage_sheet = workbook["充放电量"]
        emergency_sheet = workbook["紧急购电量"]

        output_results = [item for item in results if item.day >= date(2025, 2, 1)]
        if len(output_results) != 334:
            raise RuntimeError("结果应包含2025-02-01至2025-12-31共334天。")

        template_dates = [
            plan_sheet.cell(row, 1).value.date()
            if isinstance(plan_sheet.cell(row, 1).value, datetime)
            else plan_sheet.cell(row, 1).value
            for row in range(2, plan_sheet.max_row + 1)
        ]
        expected_dates = [item.day for item in output_results]
        if template_dates != expected_dates:
            raise ValueError("result2模板的计划购电日期与计算日期不一致。")

        for row, item in enumerate(output_results, start=2):
            grid_template = _to_template_order(item.grid_kwh, chronological_order)
            for column, value in enumerate(grid_template, start=2):
                plan_sheet.cell(row, column, round(float(value), 4))
            plan_sheet.cell(row, 146, round(float(np.sum(item.grid_kwh)), 4))
            plan_sheet.cell(row, 147, round(item.plan_cost_yuan, 4))

        storage_styles = [
            [copy(storage_sheet.cell(row, column)) for column in range(1, 7)]
            for row in (2, 3)
        ]
        if storage_sheet.max_row > 1:
            storage_sheet.delete_rows(2, storage_sheet.max_row - 1)
        storage_row = 2
        periods = [
            "0:00-4:00",
            "4:00-8:00",
            "8:00-12:00",
            "12:00-16:00",
            "16:00-20:00",
            "20:00-24:00",
        ]
        for item in output_results:
            for block, period in enumerate(periods):
                style_row = storage_styles[0 if block == 0 else 1]
                for column in range(1, 7):
                    _copy_cell_style(style_row[column - 1], storage_sheet.cell(storage_row, column))
                start, stop = block * 24, (block + 1) * 24
                storage_sheet.cell(storage_row, 1, item.day if block == 0 else None)
                storage_sheet.cell(storage_row, 2, period)
                storage_sheet.cell(
                    storage_row, 3, round(float(np.sum(item.charge_kwh[start:stop])), 4)
                )
                storage_sheet.cell(
                    storage_row,
                    4,
                    round(float(np.sum(item.discharge_kwh[start:stop])), 4),
                )
                if block == 0:
                    storage_sheet.cell(storage_row, 5, "0:00")
                    storage_sheet.cell(storage_row, 6, round(item.soc_start_kwh, 4))
                elif block == 1:
                    storage_sheet.cell(storage_row, 5, "24:00")
                    storage_sheet.cell(storage_row, 6, round(item.soc_end_kwh, 4))
                storage_row += 1

        emergency_styles = [
            [copy(emergency_sheet.cell(row, column)) for column in range(1, 4)]
            for row in (2, 3)
        ]
        if emergency_sheet.max_row > 1:
            emergency_sheet.delete_rows(2, emergency_sheet.max_row - 1)
        emergency_row = 2
        for item in output_results:
            groups = group_emergency(item.emergency_kwh, cfg.emergency_tolerance_kwh)
            for group_index, (period, amount) in enumerate(groups):
                style_row = emergency_styles[0 if group_index == 0 else 1]
                for column in range(1, 4):
                    _copy_cell_style(style_row[column - 1], emergency_sheet.cell(emergency_row, column))
                emergency_sheet.cell(
                    emergency_row, 1, item.day if group_index == 0 else None
                )
                emergency_sheet.cell(emergency_row, 2, period)
                emergency_sheet.cell(emergency_row, 3, round(amount, 4))
                emergency_row += 1
        if emergency_row == 2:
            for column in range(1, 4):
                _copy_cell_style(emergency_styles[0][column - 1], emergency_sheet.cell(2, column))
            emergency_sheet.cell(2, 1, "无紧急购电")
            emergency_sheet.cell(2, 3, 0.0)

        workbook.save(temporary)
        try:
            os.replace(temporary, output)
        except PermissionError as exc:
            raise PermissionError(
                f"无法写入{output}，请关闭正在打开该文件的Excel窗口后重试。"
            ) from exc
    finally:
        if temporary.exists():
            temporary.unlink()


def write_daily_summary(
    path: Path, historical: HistoricalData, results: list[DailyResult]
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    headers = [
        "日期",
        "计划购电量(kWh)",
        "计划购电费(元)",
        "紧急购电量(kWh)",
        "紧急购电费(元)",
        "总购电费(元)",
        "计划弃光量(kWh)",
        "真实弃光量(kWh)",
        "安全裕度合计(kWh)",
        "0:00储电量(kWh)",
        "24:00储电量(kWh)",
        "负荷预测MAE(kW)",
        "光伏预测MAE(kW)",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file, lineterminator="\n")
        writer.writerow(headers)
        for index, item in enumerate(results):
            writer.writerow(
                [
                    item.day.isoformat(),
                    f"{np.sum(item.grid_kwh):.6f}",
                    f"{item.plan_cost_yuan:.6f}",
                    f"{np.sum(item.emergency_kwh):.6f}",
                    f"{item.emergency_cost_yuan:.6f}",
                    f"{item.total_cost_yuan:.6f}",
                    f"{np.sum(item.planned_curtailment_kwh):.6f}",
                    f"{np.sum(item.actual_curtailment_kwh):.6f}",
                    f"{np.sum(item.reserve_kwh):.6f}",
                    f"{item.soc_start_kwh:.6f}",
                    f"{item.soc_end_kwh:.6f}",
                    f"{np.mean(np.abs(historical.load_kw[index] - item.forecast_load_kw)):.6f}",
                    f"{np.mean(np.abs(historical.pv_kw[index] - item.forecast_pv_kw)):.6f}",
                ]
            )


def _r_squared(actual: np.ndarray, forecast: np.ndarray) -> float:
    residual = float(np.sum((actual - forecast) ** 2))
    total = float(np.sum((actual - np.mean(actual)) ** 2))
    return float("nan") if total <= 0.0 else 1.0 - residual / total


def _split_indices(historical: HistoricalData, split: DatasetSplit) -> dict[str, list[int]]:
    groups = split.cycle_labels(historical.dates)
    if any(not indices for indices in groups.values()):
        raise ValueError(f"训练/验证/测试划分产生空数据集：{groups}")
    return groups


def split_period_metrics(
    historical: HistoricalData,
    results: list[DailyResult],
    indices: list[int],
) -> dict[str, float | int | str]:
    selected = [results[index] for index in indices]
    actual_load = historical.load_kw[indices]
    actual_pv = historical.pv_kw[indices]
    forecast_load = np.asarray([item.forecast_load_kw for item in selected])
    forecast_pv = np.asarray([item.forecast_pv_kw for item in selected])
    actual_net = actual_load - actual_pv
    forecast_net = forecast_load - forecast_pv
    plan_energy = float(sum(np.sum(item.grid_kwh) for item in selected))
    emergency_energy = float(sum(np.sum(item.emergency_kwh) for item in selected))
    plan_cost = float(sum(item.plan_cost_yuan for item in selected))
    emergency_cost = float(sum(item.emergency_cost_yuan for item in selected))
    purchased = plan_energy + emergency_energy
    return {
        "开始日期": selected[0].day.isoformat(),
        "结束日期": selected[-1].day.isoformat(),
        "天数": len(selected),
        "负荷MAE(kW)": float(np.mean(np.abs(actual_load - forecast_load))),
        "负荷RMSE(kW)": float(np.sqrt(np.mean((actual_load - forecast_load) ** 2))),
        "负荷R2": _r_squared(actual_load, forecast_load),
        "光伏MAE(kW)": float(np.mean(np.abs(actual_pv - forecast_pv))),
        "光伏RMSE(kW)": float(np.sqrt(np.mean((actual_pv - forecast_pv) ** 2))),
        "光伏R2": _r_squared(actual_pv, forecast_pv),
        "净负荷MAE(kW)": float(np.mean(np.abs(actual_net - forecast_net))),
        "净负荷RMSE(kW)": float(np.sqrt(np.mean((actual_net - forecast_net) ** 2))),
        "计划购电量(kWh)": plan_energy,
        "紧急购电量(kWh)": emergency_energy,
        "紧急购电占比": 0.0 if purchased <= 0.0 else emergency_energy / purchased,
        "计划购电费(元)": plan_cost,
        "紧急购电费(元)": emergency_cost,
        "总购电费(元)": plan_cost + emergency_cost,
    }


def write_split_evaluation(
    path: Path,
    historical: HistoricalData,
    results: list[DailyResult],
    split: DatasetSplit,
) -> None:
    groups = _split_indices(historical, split)
    rows = []
    for label in ("训练集", "验证集", "测试集"):
        row = {"数据集": label}
        row.update(split_period_metrics(historical, results, groups[label]))
        rows.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    headers = list(rows[0].keys())
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=headers, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)

    print("\n问题二训练/验证/测试检验")
    for row in rows:
        print(
            f"  {row['数据集']} {row['开始日期']}至{row['结束日期']}："
            f"净负荷RMSE={row['净负荷RMSE(kW)']:.4f} kW，"
            f"总费用={row['总购电费(元)']:.4f} 元，"
            f"紧急购电占比={row['紧急购电占比']:.4%}"
        )
    print(f"数据集检验报告：{path.resolve()}")


def validate_year(
    historical: HistoricalData,
    results: list[DailyResult],
    cfg: Config,
) -> None:
    if len(results) != len(historical.dates):
        raise RuntimeError("全年结果数量与历史日期数量不一致。")
    tolerance = 1.0e-4
    if abs(results[0].soc_start_kwh - cfg.initial_soc_kwh) > tolerance:
        raise RuntimeError("2025-01-01初始SOC不是6000 kWh。")
    for index, item in enumerate(results):
        if item.day != historical.dates[index]:
            raise RuntimeError("逐日结果日期错位。")
        if index > 0 and abs(item.soc_start_kwh - results[index - 1].soc_end_kwh) > tolerance:
            raise RuntimeError(f"{item.day}的SOC没有与前一天连续衔接。")
        if not (cfg.soc_min_kwh - tolerance <= item.soc_end_kwh <= cfg.soc_max_kwh + tolerance):
            raise RuntimeError(f"{item.day}的日末SOC越界。")
        if np.any(item.grid_kwh < -tolerance) or np.any(item.emergency_kwh < -tolerance):
            raise RuntimeError(f"{item.day}存在负购电量。")


def print_summary(
    historical: HistoricalData,
    results: list[DailyResult],
    output: Path,
    summary: Path,
) -> None:
    output_results = [item for item in results if item.day >= date(2025, 2, 1)]
    plan_energy = float(sum(np.sum(item.grid_kwh) for item in output_results))
    emergency_energy = float(sum(np.sum(item.emergency_kwh) for item in output_results))
    plan_cost = float(sum(item.plan_cost_yuan for item in output_results))
    emergency_cost = float(sum(item.emergency_cost_yuan for item in output_results))
    emergency_intervals = int(
        sum(np.count_nonzero(item.emergency_kwh > 1.0e-5) for item in output_results)
    )

    start_index = historical.dates.index(date(2025, 2, 1))
    actual_load = historical.load_kw[start_index:]
    actual_pv = historical.pv_kw[start_index:]
    forecast_load = np.asarray([item.forecast_load_kw for item in output_results])
    forecast_pv = np.asarray([item.forecast_pv_kw for item in output_results])
    load_mae = float(np.mean(np.abs(actual_load - forecast_load)))
    load_rmse = float(np.sqrt(np.mean((actual_load - forecast_load) ** 2)))
    pv_mae = float(np.mean(np.abs(actual_pv - forecast_pv)))
    pv_rmse = float(np.sqrt(np.mean((actual_pv - forecast_pv) ** 2)))

    print("\n问题二计算结果（2025-02-01至2025-12-31）")
    print(f"  计划购电量：{plan_energy:.4f} kWh")
    print(f"  计划购电费：{plan_cost:.4f} 元")
    print(f"  紧急购电量：{emergency_energy:.4f} kWh")
    print(f"  紧急购电费：{emergency_cost:.4f} 元")
    print(f"  总购电费：{plan_cost + emergency_cost:.4f} 元")
    print(f"  发生紧急购电的10分钟时段数：{emergency_intervals}")
    print(f"  负荷滚动预测：MAE={load_mae:.4f} kW，RMSE={load_rmse:.4f} kW")
    print(f"  光伏滚动预测：MAE={pv_mae:.4f} kW，RMSE={pv_rmse:.4f} kW")
    print(
        f"  SOC范围：{min(item.soc_start_kwh for item in results):.4f}～"
        f"{max(item.soc_end_kwh for item in results):.4f} kWh"
    )

    specified = {date(2025, 3, 20), date(2025, 6, 21), date(2025, 9, 23), date(2025, 12, 21)}
    print("\n表3指定日期紧急购电")
    for item in output_results:
        if item.day not in specified:
            continue
        groups = group_emergency(item.emergency_kwh, 1.0e-5)
        if not groups:
            print(f"  {item.day.isoformat()}：无")
        else:
            print(f"  {item.day.isoformat()}：")
            for period, amount in groups:
                print(f"    {period}：{amount:.4f} kWh")

    print(f"\n正式结果：{output.resolve()}")
    print(f"逐日汇总：{summary.resolve()}")


@dataclass(frozen=True)
class LoadSearchPoint:
    """负荷预测贝叶斯优化的一个超参数候选。"""

    history_days: int
    half_life_days: float
    same_weekday_multiplier: float
    wavelet_level: int
    annual_harmonics: int
    load_ridge_lambda: float
    dynamic_load_blend: float

    def config(self) -> Config:
        return replace(Config(), **asdict(self))


def _decode_load_point(unit: np.ndarray) -> LoadSearchPoint:
    x = np.clip(np.asarray(unit, dtype=float), 0.0, 1.0)
    return LoadSearchPoint(
        history_days=int(round(14 + x[0] * 46)),
        half_life_days=float(3.0 + x[1] * 27.0),
        same_weekday_multiplier=float(x[2] * 5.0),
        wavelet_level=int(round(1 + x[3] * 3)),
        annual_harmonics=int(round(1 + x[4] * 7)),
        load_ridge_lambda=float(10.0 ** (-4.0 + 7.0 * x[5])),
        dynamic_load_blend=float(0.20 + x[6] * 0.80),
    )


def _encode_load_point(point: LoadSearchPoint) -> np.ndarray:
    return np.asarray(
        [
            (point.history_days - 14) / 46,
            (point.half_life_days - 3.0) / 27.0,
            point.same_weekday_multiplier / 5.0,
            (point.wavelet_level - 1) / 3.0,
            (point.annual_harmonics - 1) / 7.0,
            (math.log10(point.load_ridge_lambda) + 4.0) / 7.0,
            (point.dynamic_load_blend - 0.20) / 0.80,
        ],
        dtype=float,
    )


def _load_point_key(point: LoadSearchPoint) -> tuple[object, ...]:
    return (
        point.history_days,
        round(point.half_life_days, 6),
        round(point.same_weekday_multiplier, 6),
        point.wavelet_level,
        point.annual_harmonics,
        round(math.log10(point.load_ridge_lambda), 6),
        round(point.dynamic_load_blend, 6),
    )


def _matern52(
    a: np.ndarray, b: np.ndarray, length_scale: float
) -> np.ndarray:
    distance = np.sqrt(
        np.maximum(0.0, np.sum((a[:, None, :] - b[None, :, :]) ** 2, axis=2))
    )
    scaled = math.sqrt(5.0) * distance / length_scale
    return (1.0 + scaled + scaled**2 / 3.0) * np.exp(-scaled)


def _expected_improvement(
    observed_x: np.ndarray,
    observed_y: np.ndarray,
    candidates: np.ndarray,
    length_scale: float,
) -> np.ndarray:
    mean_y = float(np.mean(observed_y))
    scale_y = max(float(np.std(observed_y)), 1.0e-12)
    normalized_y = (observed_y - mean_y) / scale_y
    kernel = _matern52(observed_x, observed_x, length_scale)
    kernel.flat[:: kernel.shape[0] + 1] += 1.0e-7
    try:
        chol = np.linalg.cholesky(kernel)
    except np.linalg.LinAlgError:
        kernel.flat[:: kernel.shape[0] + 1] += 1.0e-5
        chol = np.linalg.cholesky(kernel)
    alpha = np.linalg.solve(chol.T, np.linalg.solve(chol, normalized_y))
    cross = _matern52(observed_x, candidates, length_scale)
    mean = cross.T @ alpha
    solved = np.linalg.solve(chol, cross)
    sigma = np.sqrt(np.maximum(1.0e-12, 1.0 - np.sum(solved**2, axis=0)))
    best = float(np.min(normalized_y))
    improvement = best - mean - 0.001
    z = improvement / sigma
    cdf = 0.5 * (
        1.0 + np.asarray([math.erf(value / math.sqrt(2.0)) for value in z])
    )
    pdf = np.exp(-0.5 * z**2) / math.sqrt(2.0 * math.pi)
    return improvement * cdf + sigma * pdf


def _predict_load_indices(
    baseline: BaselineData,
    historical: HistoricalData,
    cfg: Config,
    indices: list[int],
    training_indices: set[int],
    incumbent_rmse: float | None = None,
    early_stop_margin: float = 0.01,
) -> tuple[np.ndarray | None, float, bool]:
    hours = np.arange(144, dtype=float) / 6.0
    forecasts: list[np.ndarray] = []
    squared_error = 0.0
    total_values = len(indices) * 144
    for position, index in enumerate(indices, start=1):
        forecast = make_forecast(
            historical,
            baseline,
            index,
            historical.dates[index],
            hours,
            cfg,
            allowed_indices=training_indices,
        ).load_kw
        forecasts.append(forecast)
        squared_error += float(np.sum((historical.load_kw[index] - forecast) ** 2))
        if incumbent_rmse is not None and position >= 8:
            lower_bound = math.sqrt(squared_error / total_values)
            if lower_bound > incumbent_rmse * (1.0 + early_stop_margin):
                return None, lower_bound, True
    return (
        np.asarray(forecasts, dtype=float),
        math.sqrt(squared_error / total_values),
        False,
    )


def _load_metrics(actual: np.ndarray, forecast: np.ndarray) -> dict[str, float]:
    error = actual - forecast
    return {
        "mae_kw": float(np.mean(np.abs(error))),
        "rmse_kw": float(np.sqrt(np.mean(error**2))),
        "r2": _r_squared(actual, forecast),
    }


def run_load_bayesian_optimization(args: argparse.Namespace) -> None:
    """严格前向寻优；仅写参数评价报告，不执行购电线性规划。"""
    if args.bayes_trials < 2 or not 1 <= args.bayes_initial < args.bayes_trials:
        raise ValueError("应满足 bayes-trials >= 2 且 1 <= bayes-initial < bayes-trials。")
    baseline, order = read_baseline(args.baseline)
    historical = read_historical_data(args.load, args.pv, order)
    groups = DatasetSplit().cycle_labels(historical.dates)
    training_indices = set(groups["训练集"])
    labels = {
        index: label for label, indices in groups.items() for index in indices
    }
    rng = np.random.default_rng(args.bayes_seed)
    base_cfg = Config()
    default_point = LoadSearchPoint(
        base_cfg.history_days,
        base_cfg.half_life_days,
        base_cfg.same_weekday_multiplier,
        base_cfg.wavelet_level,
        base_cfg.annual_harmonics,
        base_cfg.load_ridge_lambda,
        base_cfg.dynamic_load_blend,
    )
    current = default_point
    candidates = [default_point]
    candidate_sse = [0.0]
    seen = {_load_point_key(default_point)}
    validation_seen: list[int] = []
    all_forecasts: list[np.ndarray] = []
    proposals: list[dict[str, object]] = []
    selection_history: list[dict[str, object]] = []
    proposal_count = 1
    no_improvement = 0
    search_stopped = False
    hours = np.arange(144, dtype=float) / 6.0

    print("严格前向负荷贝叶斯优化：测试集不参与参数搜索。")
    for index, day in enumerate(historical.dates):
        current_forecast = make_forecast(
            historical,
            baseline,
            index,
            day,
            hours,
            current.config(),
            allowed_indices=training_indices,
        ).load_kw
        all_forecasts.append(current_forecast)
        if labels[index] != "验证集":
            continue
        validation_seen.append(index)
        actual = historical.load_kw[index]
        for candidate_index, point in enumerate(candidates):
            forecast = (
                current_forecast
                if _load_point_key(point) == _load_point_key(current)
                else make_forecast(
                    historical,
                    baseline,
                    index,
                    day,
                    hours,
                    point.config(),
                    allowed_indices=training_indices,
                ).load_kw
            )
            candidate_sse[candidate_index] += float(np.sum((actual - forecast) ** 2))

        denominator = len(validation_seen) * 144
        scores = np.sqrt(np.asarray(candidate_sse) / denominator)
        incumbent = float(np.min(scores))
        if proposal_count < args.bayes_trials and not search_stopped:
            if len(candidates) < args.bayes_initial:
                unit = rng.random(7)
                method = "随机初始化"
            else:
                pool = rng.random((args.bayes_candidate_pool, 7))
                acquisition = _expected_improvement(
                    np.asarray([_encode_load_point(point) for point in candidates]),
                    scores,
                    pool,
                    0.35,
                )
                unit = pool[int(np.argmax(acquisition))]
                method = "贝叶斯期望改进"
            point = _decode_load_point(unit)
            attempts = 0
            while _load_point_key(point) in seen and attempts < 100:
                point = _decode_load_point(rng.random(7))
                attempts += 1
            seen.add(_load_point_key(point))
            proposal_count += 1
            _, proposed_score, candidate_stopped = _predict_load_indices(
                baseline,
                historical,
                point.config(),
                validation_seen,
                training_indices,
                incumbent,
            )
            improved = (
                not candidate_stopped
                and proposed_score < incumbent - args.bayes_min_delta
            )
            if not candidate_stopped:
                candidates.append(point)
                candidate_sse.append(proposed_score**2 * denominator)
                no_improvement = 0 if improved else no_improvement + int(
                    len(candidates) >= args.bayes_initial
                )
            elif len(candidates) >= args.bayes_initial:
                no_improvement += 1
            proposals.append(
                {
                    "proposal": proposal_count,
                    "validation_date": day.isoformat(),
                    "validation_days_available": len(validation_seen),
                    "method": method,
                    "validation_rmse_kw": proposed_score,
                    "candidate_early_stopped": candidate_stopped,
                    "improved": improved,
                    "parameters": asdict(point),
                }
            )
            print(f"[{day}] 累计验证RMSE={proposed_score:.4f} kW")
            if len(candidates) >= args.bayes_initial and no_improvement >= args.bayes_patience:
                search_stopped = True
                print(f"连续{args.bayes_patience}个候选无显著改善，停止寻优。")

        scores = np.sqrt(np.asarray(candidate_sse) / denominator)
        selected_index = int(np.argmin(scores))
        if len(validation_seen) >= args.bayes_initial:
            current = candidates[selected_index]
        selection_history.append(
            {
                "date": day.isoformat(),
                "validation_days_available": len(validation_seen),
                "selected_validation_rmse_kw": float(scores[selected_index]),
                "selected_parameters": asdict(current),
            }
        )

    forward = np.asarray(all_forecasts, dtype=float)
    final_metrics = {
        label: _load_metrics(historical.load_kw[indices], forward[indices])
        for label, indices in groups.items()
    }
    baseline_metrics: dict[str, dict[str, float]] = {}
    for label, indices in groups.items():
        forecast, _, stopped = _predict_load_indices(
            baseline,
            historical,
            default_point.config(),
            indices,
            training_indices,
        )
        if stopped or forecast is None:
            raise RuntimeError("最终评价不应提前停止。")
        baseline_metrics[label] = _load_metrics(historical.load_kw[indices], forecast)
    report = {
        "data_split": {
            "rule": "训练、训练、训练、验证、训练、测试",
            "counts": {label: len(indices) for label, indices in groups.items()},
            "strict_training_freeze": True,
            "test_data_used_in_optimization": False,
        },
        "optimization": {
            "method": "strict forward Gaussian-process Bayesian optimization",
            "objective": "截至当前日期的累计验证集RMSE",
            "early_stopping_patience": args.bayes_patience,
            "minimum_improvement_kw": args.bayes_min_delta,
            "maximum_proposals": args.bayes_trials,
            "completed_proposals": proposal_count,
            "search_early_stopped": search_stopped,
            "seed": args.bayes_seed,
        },
        "final_selected_parameters": asdict(current),
        "baseline_metrics": baseline_metrics,
        "metrics": final_metrics,
        "proposals": proposals,
        "selection_history": selection_history,
    }
    args.bayes_report.parent.mkdir(parents=True, exist_ok=True)
    args.bayes_report.write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    for label in ("训练集", "验证集", "测试集"):
        values = final_metrics[label]
        print(
            f"{label}：MAE={values['mae_kw']:.4f} kW，"
            f"RMSE={values['rmse_kw']:.4f} kW，R2={values['r2']:.6f}"
        )
    print(f"实验报告：{args.bayes_report.resolve()}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="C题问题二滚动预测与购电线性规划")
    parser.add_argument("--baseline", type=Path, default=DEFAULT_BASELINE)
    parser.add_argument("--load", type=Path, default=DEFAULT_LOAD)
    parser.add_argument("--pv", type=Path, default=DEFAULT_PV)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--summary", type=Path, default=DEFAULT_SUMMARY)
    parser.add_argument("--split-report", type=Path, default=DEFAULT_SPLIT_REPORT)
    parser.add_argument("--train-end", type=date.fromisoformat, default=date(2025, 7, 31))
    parser.add_argument(
        "--validation-end", type=date.fromisoformat, default=date(2025, 9, 30)
    )
    parser.add_argument(
        "--bayes-optimize",
        action="store_true",
        help="执行严格前向贝叶斯参数寻优，不生成正式购电结果表。",
    )
    parser.add_argument("--bayes-trials", type=int, default=40)
    parser.add_argument("--bayes-initial", type=int, default=10)
    parser.add_argument("--bayes-patience", type=int, default=12)
    parser.add_argument("--bayes-min-delta", type=float, default=0.05)
    parser.add_argument("--bayes-candidate-pool", type=int, default=2000)
    parser.add_argument("--bayes-seed", type=int, default=2026)
    parser.add_argument("--bayes-report", type=Path, default=DEFAULT_BAYES_REPORT)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if args.bayes_optimize:
        run_load_bayesian_optimization(args)
        return
    cfg = Config()
    baseline, chronological_order = read_baseline(args.baseline)
    historical = read_historical_data(args.load, args.pv, chronological_order)
    split = DatasetSplit(args.train_end, args.validation_end)
    groups = split.cycle_labels(historical.dates)
    training_indices = set(groups["训练集"])
    results = simulate_year(
        baseline, historical, cfg, training_indices=training_indices
    )
    validate_year(historical, results, cfg)
    write_result_workbook(
        args.template,
        args.output,
        results,
        baseline.price,
        chronological_order,
        cfg,
    )
    write_daily_summary(args.summary, historical, results)
    write_split_evaluation(args.split_report, historical, results, split)
    print_summary(historical, results, args.output, args.summary)


if __name__ == "__main__":
    main()
